from pathlib import Path
import json
from docx import Document
from docx.oxml.ns import qn

SOURCE = Path(r"C:\Python\tesis\documentacion\TESIS_AGO2026_Rev57_(ZUJ).docx")
OUT = Path(r"C:\Python\tesis\tmp\auditoria_rev57\estructura.json")

doc = Document(SOURCE)
items = []
page = 1

for idx, p in enumerate(doc.paragraphs, start=1):
    text = p.text.strip()
    breaks_before = len(p._p.xpath('.//w:lastRenderedPageBreak'))
    page += breaks_before
    if text or breaks_before:
        items.append({
            "kind": "paragraph",
            "index": idx,
            "page": page,
            "style": p.style.name if p.style else "",
            "text": text,
            "page_breaks_before": breaks_before,
        })

tables = []
for ti, table in enumerate(doc.tables, start=1):
    rows = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    tables.append({"table": ti, "rows": rows})

payload = {
    "source": str(SOURCE),
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "last_rendered_page": page,
    "paragraphs": items,
    "tables": tables,
}
OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({k: payload[k] for k in ("paragraph_count", "table_count", "last_rendered_page")}, ensure_ascii=False))
