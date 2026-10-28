from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

SOURCE = Path(r"C:\Python\tesis\documentacion\TESIS_OCT2026_Rev67_(ZUJ)_Final.docx")
OUTPUT = Path(r"C:\Python\tesis\documentacion\TESIS_OCT2026_Rev68_(ZUJ)_Final.docx")


def iter_all_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def replace_preserving_runs(paragraph, old, new):
    """Replace text in all w:t nodes, including text nested in hyperlinks/fields."""
    count = 0
    while True:
        nodes = list(paragraph._p.iter(qn("w:t")))
        full = "".join(node.text or "" for node in nodes)
        start = full.find(old)
        if start < 0:
            break
        end = start + len(old)
        positions = []
        cursor = 0
        for idx, node in enumerate(nodes):
            node_text = node.text or ""
            positions.append((idx, cursor, cursor + len(node_text)))
            cursor += len(node_text)
        first = next(idx for idx, a, b in positions if a <= start < b or (a == b == start))
        last = next(idx for idx, a, b in positions if a < end <= b)
        first_a = positions[first][1]
        last_a = positions[last][1]
        prefix = (nodes[first].text or "")[: start - first_a]
        suffix = (nodes[last].text or "")[end - last_a :]
        nodes[first].text = prefix + new + suffix
        for idx in range(first + 1, last + 1):
            nodes[idx].text = ""
        count += 1
    return count


def set_paragraph_text_preserving_format(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


doc = Document(SOURCE)
counts = {}

replacements = [
    ("Ben Abdallah et al. (2022)", "Omrane (2022)"),
    ("(Ben Abdallah et al., 2022)", "(Omrane, 2022)"),
    ("Dubey et al. (2021)", "Dubey et al. (2020)"),
    ("(Dubey et al., 2021)", "(Dubey et al., 2020)"),
    ("Saavedra García y Tapia Sánchez (2020)", "Saavedra García y Tapia Sánchez (2013)"),
    ("Saavedra-García y Tapia-Sánchez (2020)", "Saavedra-García y Tapia-Sánchez (2013)"),
    ("(Makridakis et al., 2018; Spiliotis et al., 2022)", "(Makridakis et al., 2018a; Spiliotis et al., 2022)"),
    ("(Makridakis et al., 2018, 2020)", "(Makridakis et al., 2018b, 2020)"),
    ("Makridakis et al. (2018)", "Makridakis et al. (2018b)"),
    ("; Spiliotis et al., 2025", ""),
]

for old, new in replacements:
    total = 0
    for paragraph in iter_all_paragraphs(doc):
        total += replace_preserving_runs(paragraph, old, new)
    counts[old] = total

reference_updates = {
    "Ben Abdallah, H., Ayadi, M., & Boujelbene, Y. (2022).": (
        "Omrane, A. (2022). The main determinants and effects of product innovation: "
        "An exploratory study on the pastry companies of the region of Sfax (in Tunisia). "
        "Technological Forecasting and Social Change, 185, Article 122065. "
        "https://doi.org/10.1016/j.techfore.2022.122065"
    ),
    "Dubey, R., Gunasekaran, A., Childe, S. J., Bryde, D. J., Giannakis, M., & Foropon, C. (2021).": (
        "Dubey, R., Gunasekaran, A., Childe, S. J., Bryde, D. J., Giannakis, M., & Foropon, C. (2020). "
        "Big data analytics and artificial intelligence pathway to operational performance under the effects of "
        "entrepreneurial orientation and environmental dynamism: A study of manufacturing organisations. "
        "International Journal of Production Economics, 226, 107599. "
        "https://doi.org/10.1016/j.ijpe.2019.107599"
    ),
    "Saavedra-García, M., & Tapia-Sánchez, B. (2020).": (
        "Saavedra García, M. L., & Tapia Sánchez, B. (2013). El uso de las tecnologías de información y comunicación "
        "TIC en las micro, pequeñas y medianas empresas (MIPyME) industriales mexicanas. Enl@ce: Revista Venezolana "
        "de Información, Tecnología y Conocimiento, 10(1), 85–104. "
        "https://www.redalyc.org/articulo.oa?id=82326270007"
    ),
}

for marker, new_text in reference_updates.items():
    matches = [p for p in doc.paragraphs if p.text.startswith(marker)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one reference starting with {marker!r}, found {len(matches)}")
    set_paragraph_text_preserving_format(matches[0], new_text)
    counts[f"reference:{marker}"] = 1

spiliotis_marker = "Spiliotis, E., Makridakis, S., & Assimakopoulos, V. (2025)."
spiliotis_matches = [p for p in doc.paragraphs if p.text.startswith(spiliotis_marker)]
if len(spiliotis_matches) != 1:
    raise RuntimeError(f"Expected one invalid REF-40 paragraph, found {len(spiliotis_matches)}")
remove_paragraph(spiliotis_matches[0])
counts["removed_REF-40"] = 1

# Request Word to refresh page numbers, TOC and other fields when opened.
settings = doc.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = settings.makeelement(qn("w:updateFields"), {qn("w:val"): "true"})
    settings.append(update_fields)
else:
    update_fields.set(qn("w:val"), "true")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)

# Structural validation on the saved file.
check = Document(OUTPUT)
all_text = "\n".join(p.text for p in iter_all_paragraphs(check))
for forbidden in [
    "Ben Abdallah et al. (2022)", "Ben Abdallah et al., 2022", "Dubey et al. (2021)", "Dubey et al., 2021",
    "Saavedra García y Tapia Sánchez (2020)", "Saavedra-García y Tapia-Sánchez (2020)",
    "Spiliotis et al., 2025", "Continuous business forecasting in modern organizations",
    "Makridakis et al. (2018)", "Makridakis et al., 2018;", "Makridakis et al., 2018, 2020",
]:
    if forbidden in all_text:
        raise RuntimeError(f"Forbidden legacy text remains: {forbidden}")

required = {
    "Omrane (2022)": 1,
    "Omrane, 2022": 1,
    "Dubey et al. (2020)": 3,
    "Dubey et al., 2020": 1,
    "Saavedra García y Tapia Sánchez (2013)": 1,
    "Saavedra-García y Tapia-Sánchez (2013)": 1,
    "Makridakis et al., 2018a": 3,
    "Makridakis et al., 2018b": 3,
    "Makridakis et al. (2018b)": 1,
    "Omrane, A. (2022).": 1,
    "Dubey, R., Gunasekaran, A., Childe, S. J., Bryde, D. J., Giannakis, M., & Foropon, C. (2020).": 1,
    "Saavedra García, M. L., & Tapia Sánchez, B. (2013).": 1,
}
for needle, minimum in required.items():
    actual = all_text.count(needle)
    if actual < minimum:
        raise RuntimeError(f"Required text {needle!r}: expected at least {minimum}, found {actual}")

print({"output": str(OUTPUT), "replacements": counts, "paragraphs": len(check.paragraphs), "tables": len(check.tables)})
