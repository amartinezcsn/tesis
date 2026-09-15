"""Read-only source audit. Prints aggregated evidence; never fills gaps."""
from openpyxl import load_workbook
from datetime import datetime,timedelta
from collections import Counter,defaultdict
from pathlib import Path
import hashlib,json,math

start=datetime(2021,5,1);end=datetime(2024,6,30)
def date(v):
    if isinstance(v,datetime):return v.replace(hour=0,minute=0,second=0,microsecond=0)
    if isinstance(v,str):
        for fmt in ('%Y-%m-%d','%Y-%m-%d %H:%M:%S','%d/%m/%Y'):
            try:return datetime.strptime(v,fmt)
            except ValueError:pass
    return None
def num(v):
    try:
        n=float(v);return n if math.isfinite(n) else None
    except (ValueError,TypeError):return None
def monday(d):return d-timedelta(days=d.weekday())
def runs(weeks):
    result=[]
    for w in sorted(weeks):
        if result and w==result[-1][-1]+timedelta(days=7):result[-1].append(w)
        else:result.append([w])
    return [{'inicio':str(r[0].date()),'fin_lunes':str(r[-1].date()),'semanas':len(r)} for r in result]

path=Path('datasets/xlsx/Compras.xlsx');sha=hashlib.sha256(path.read_bytes()).hexdigest()
w=load_workbook(path,read_only=True,data_only=True);rows=list(w.active.values);w.close()
headers=rows[0];allrows=rows[1:];period=[(i+2,r) for i,r in enumerate(allrows) if date(r[1]) and start<=date(r[1])<=end]
dups=Counter(tuple(r) for r in allrows)
exclude={'COMBUSTIBLE','MUEBLES','HERRAMIENTA'}
uncertain={'CONSUMIBLES','OTROS','NO DEFINIDO','CASETAS','ADMINISTRACION','HORNEADO'}
weeks=[datetime(2021,5,3)+timedelta(weeks=i) for i in range(165)]
assert weeks[-1]==datetime(2024,6,24)
counts=Counter();included=Counter();categories=defaultdict(lambda:[0,0]);issues=[];monthly=Counter();positive=Counter();zero=0
amb=defaultdict(set)
for i,r in period:
    d=date(r[1]);week=monday(d);amount=num(r[6]);category=str(r[8] or '').strip().upper()
    counts[week]+=1;monthly[d.strftime('%Y-%m')]+=1
    scope='excluir' if category in exclude else 'revisar' if category in uncertain or not category else 'candidata_inclusion'
    categories[(category,scope)][0]+=1
    if amount is not None:categories[(category,scope)][1]+=amount
    if amount==0:zero+=1
    if amount is not None and amount>0:positive[week]+=1
    if scope=='revisar':amb[category].add(str(r[5]))
    if amount is None or amount<0 or dups[tuple(r)]>1 or not str(r[5] or '').strip():
        issues.append({'fila':i,'fecha':str(d.date()),'categoria':category,'importe':amount,'duplicado':dups[tuple(r)]>1})
    if scope=='candidata_inclusion' and amount is not None and amount>0:included[week]+=1
missing=[x for x in weeks if not counts[x]]
out={'source':str(path),'sha256':sha,'periodo':['2021-05-01','2024-06-30'],'registros_periodo':len(period),
'fechas_min_max':[str(min(date(r[1]) for _,r in period).date()),str(max(date(r[1]) for _,r in period).date())],
'semanas_completas':len(weeks),'semanas_con_filas':sum(bool(counts[x]) for x in weeks),'semanas_sin_filas':len(missing),
'registros_cero':zero,'semanas_con_filas_sin_importe_positivo':[str(x.date()) for x in weeks if counts[x] and not positive[x]],
'huecos':runs(missing),'bloques_con_registros':sorted(runs([x for x in weeks if counts[x]]),key=lambda x:x['semanas'],reverse=True),
'semanas_con_candidatos_incluidos':sum(bool(included[x]) for x in weeks),
'bloques_candidatos_incluidos':sorted(runs([x for x in weeks if included[x]]),key=lambda x:x['semanas'],reverse=True),
'solo_excluidos_o_ambiguos':[str(x.date()) for x in weeks if counts[x] and not included[x]],
'categorias':[{'categoria':k[0],'alcance':k[1],'filas':v[0],'importe':round(v[1],2)} for k,v in sorted(categories.items())],
'incidencias':issues,'ambiguas':{k:sorted(v) for k,v in amb.items()},'meses':dict(sorted(monthly.items())),
'fechas_invalidas_fuera_o_dentro_indeterminable':[i+2 for i,r in enumerate(allrows) if not date(r[1])]}
assert hashlib.sha256(path.read_bytes()).hexdigest()==sha
print(json.dumps(out,ensure_ascii=False,indent=2))
for f in [Path('input/compras_limpias.xlsx'),Path('input/compras_limpias_corregidas.xlsx'),Path('codigos/outputs/complemento_rolling_2022_2024/Compras_complementadas_entrenamiento.xlsx')]:
    wb=load_workbook(f,read_only=True,data_only=True)
    print('COMPARACION',str(f))
    for s in wb:
        if s.title not in ('detalle','Datos'):continue
        it=iter(s.values);head=next(it);print('HOJA',s.title,'HEAD',head)
        ix=next((i for i,h in enumerate(head) if str(h).lower()=='fecha'),None)
        if ix is not None:
            ds=[date(r[ix]) for r in it];ds=[d for d in ds if d and start<=d<=end]
            print('FILAS_PERIODO',len(ds),'INICIO',min(ds) if ds else None,'FIN',max(ds) if ds else None)
    wb.close()
