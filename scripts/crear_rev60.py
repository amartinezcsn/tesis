from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Python\tesis\documentacion\TESIS_SEP2026_Rev59_(ZUJ)_Final.docx")
OUTPUT = Path(r"C:\Python\tesis\documentacion\TESIS_SEP2026_Rev60_(ZUJ)_Final.docx")
REPOSITORY_URL = "https://github.com/amartinezcsn/tesis"


UNCITED_REFERENCE_PREFIXES = (
    "Akter, S.,",
    "Bertsimas, D.,",
    "Box, G. E. P. (1976).",
    "Box, G. E. P., Jenkins,",
    "Brynjolfsson, E., & McElheran,",
    "Brynjolfsson, E., Hitt,",
    "Davenport, T. H.,",
    "Dubey, R., Bryde, D.,",
    "García-Murillo, M.,",
    "García-Pérez, A.,",
    "Hübner, A.,",
    "James, G.,",
    "Kahneman, D. (2011).",
    "Kahneman, D., & Tversky, A. (1974).",
    "Kahneman, D., & Tversky, A. (1979).",
    "Martins-Turner, K.,",
    "Makridakis, S., Spiliotis, E., & Assimakopoulos, V. (2022).",
    "Mikalef, P., Krogstie,",
    "Nonaka, I.,",
    "OECD. (2022).",
    "Ordonez Bolanos, A. A.,",
    "Petropoulos, F.,",
    "Polanyi, M. (1966).",
    "Provost, F.,",
    "Scuotto, V.,",
    "Shmueli, G., & Koppius,",
    "Simon, H. A. (1997).",
    "Sánchez, M., & Terrazas,",
    "Verhoef, P. C.,",
    "Viteri, C.,",
)


JOURNAL_NAMES = sorted(
    {
        "British Journal of Management",
        "Computational Statistics & Data Analysis",
        "Contaduría y Administración",
        "Econometrica",
        "Entrepreneurship Theory and Practice",
        "European Journal of Operational Research",
        "Foods",
        "Forecasting",
        "Humanities",
        "Industrial Marketing Management",
        "Information & Management",
        "Information Sciences",
        "Information (Switzerland)",
        "International Journal of Contemporary Hospitality Management",
        "International Journal of Forecasting",
        "International Journal of Innovation Science",
        "International Journal of Production Economics",
        "International Journal of Production Research",
        "Journal of Industrial Ecology",
        "Journal of Information & Knowledge Management",
        "Journal of the Operational Research Society",
        "Journal of the Royal Statistical Society: Series B (Methodological)",
        "Machine Learning",
        "Neurocomputing",
        "Operational Research",
        "Operational Research Quarterly",
        "PLOS ONE",
        "Procedia Computer Science",
        "Revista Colombiana De Tecnologias De Avanzada",
        "Revista Pulso Científico",
        "Statistical Science",
        "Sustainability",
        "Technological Forecasting and Social Change",
        "The Annals of Statistics",
    },
    key=len,
    reverse=True,
)


def remove_paragraph(paragraph):
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def clear_paragraph_content(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_run_font(run, name="Arial", size=8.5, bold=None, italic=None, color="000000"):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), name)


def append_styled_segments(paragraph, text, italic_spans):
    clear_paragraph_content(paragraph)
    boundaries = {0, len(text)}
    for start, end in italic_spans:
        boundaries.add(start)
        boundaries.add(end)
    points = sorted(boundaries)
    for start, end in zip(points, points[1:]):
        if start == end:
            continue
        run = paragraph.add_run(text[start:end])
        is_italic = any(a <= start and end <= b for a, b in italic_spans)
        set_run_font(run, name="Arial", size=11, italic=is_italic)


def reference_italic_spans(text):
    spans = []
    journal = next(
        (name for name in JOURNAL_NAMES if re.search(re.escape(name) + r",\s*\d+", text)),
        None,
    )
    if journal:
        start = text.index(journal)
        spans.append((start, start + len(journal)))
        volume_match = re.match(r",\s*(\d+)", text[start + len(journal) :])
        if volume_match:
            vol_start = start + len(journal) + volume_match.start(1)
            spans.append((vol_start, vol_start + len(volume_match.group(1))))
        return spans

    year_match = re.search(r"\((?:19|20)\d{2}[a-z]?[^)]*\)\.\s+", text)
    if not year_match:
        return spans
    content_start = year_match.end()
    url_match = re.search(r"\s+https?://", text)
    content_end = url_match.start() if url_match else len(text)
    body = text[content_start:content_end].rstrip()

    if " En " in body and "(pp." in body:
        en_pos = text.index(" En ", content_start) + 4
        pp_pos = text.index("(pp.", en_pos)
        editor_end = text.find(", ", en_pos, pp_pos)
        if editor_end != -1:
            title_start = editor_end + 2
            title_end = text.rfind(" ", title_start, pp_pos)
            spans.append((title_start, title_end))
        return spans

    # Books, institutional reports and preprints: italicize the work title.
    publisher_markers = [
        ". Springer", ". Wiley", ". OTexts", ". Prodem", ". OECD Publishing",
        ". Deusto", ". Quorum Books", ". Harvard Business School Press",
        ". Oxford University Press", ". O’Reilly Media", ". Free Press",
        ". Farrar, Straus and Giroux", ". Doubleday & Company",
    ]
    end = content_end
    for marker in publisher_markers:
        pos = text.find(marker, content_start, content_end)
        if pos != -1:
            end = pos
            break
    if "arXiv preprint" in text[content_start:content_end]:
        period = text.find(". arXiv preprint", content_start, content_end)
        if period != -1:
            end = period
    if end == content_end and body.endswith("."):
        end -= 1
    if end > content_start:
        spans.append((content_start, end))
    return spans


def reconcile_and_format_references(doc):
    in_references = False
    reference_paragraphs = []
    removed = []
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if text.upper() == "REFERENCIAS" and paragraph.style.name.startswith("Heading"):
            in_references = True
            continue
        if in_references and text.upper().startswith("ANEXO A"):
            break
        if not in_references or not text:
            continue
        if text.startswith(UNCITED_REFERENCE_PREFIXES):
            removed.append(text)
            remove_paragraph(paragraph)
        else:
            reference_paragraphs.append(paragraph)

    for paragraph in reference_paragraphs:
        text = paragraph.text.strip()
        append_styled_segments(paragraph, text, reference_italic_spans(text))
        fmt = paragraph.paragraph_format
        fmt.left_indent = Inches(0.5)
        fmt.first_line_indent = Inches(-0.5)
        fmt.space_after = Pt(0)
        fmt.keep_together = False
    return removed, len(reference_paragraphs)


def replace_text_in_runs(doc, replacements):
    count = 0
    for paragraph in doc.paragraphs:
        if paragraph._p.xpath(".//w:fldChar | .//w:instrText"):
            continue
        full = paragraph.text
        replacement = full
        for old, new in replacements:
            replacement = replacement.replace(old, new)
        if replacement != full:
            if paragraph.style.name.startswith("Heading"):
                for run in paragraph.runs:
                    run.text = run.text.replace("ésta", "esta")
                count += 1
                continue
            # Targeted prose corrections are safe to rebuild; none contains fields.
            clear_paragraph_content(paragraph)
            run = paragraph.add_run(replacement)
            set_run_font(run, name="Arial", size=11)
            count += 1
    return count


def fix_static_front_matter(doc):
    replacements = 0
    for text_node in doc.element.body.xpath(".//w:t"):
        if text_node.text and "objetivo de ésta" in text_node.text:
            text_node.text = text_node.text.replace("objetivo de ésta", "objetivo de esta")
            replacements += 1
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "toc 1" and paragraph.text.startswith("Anexo B Paquete de reproducibilidad"):
            for run in paragraph.runs:
                if "\t103" in run.text:
                    run.text = run.text.replace("\t103", "\t98")
                    replacements += 1
    return replacements


def add_external_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend([rpr, text_element])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def insert_repository_paragraph(doc):
    marker = next(
        p for p in doc.paragraphs if p.text.startswith("El paquete se conserva dentro de la raíz del proyecto de tesis")
    )
    new_p = OxmlElement("w:p")
    marker._p.addnext(new_p)
    paragraph = marker._parent.add_paragraph()
    paragraph._p.getparent().remove(paragraph._p)
    new_p.getparent().replace(new_p, paragraph._p)
    paragraph.style = doc.styles["Normal"]
    label = paragraph.add_run("Repositorio del proyecto: ")
    set_run_font(label, name="Arial", size=11, bold=True)
    add_external_hyperlink(paragraph, REPOSITORY_URL, REPOSITORY_URL)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C9D6", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def is_data_table(table):
    if len(table.rows) < 2:
        return False
    first = [cell.text.strip() for cell in table.rows[0].cells]
    if not first or not first[0]:
        return False
    return any(first[0].startswith(prefix) for prefix in (
        "Autor(es)", "Etapa previa", "Horizonte", "Categoría", "Campo", "Componente"
    ))


def format_data_tables(doc):
    formatted = 0
    for table in doc.tables:
        if not is_data_table(table):
            continue
        formatted += 1
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_borders(table)
        mark_header_row(table.rows[0])
        for row_index, row in enumerate(table.rows):
            fill = "D9EAF7" if row_index == 0 else ("F4F8FB" if row_index % 2 == 0 else "FFFFFF")
            for cell in row.cells:
                set_cell_shading(cell, fill)
                set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    if row_index == 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        set_run_font(run, name="Arial", size=8.5, bold=(row_index == 0))
        # Remove exact-height constraints, which can clip wrapped text.
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            for height in list(tr_pr.findall(qn("w:trHeight"))):
                tr_pr.remove(height)
    return formatted


def format_table_captions(doc):
    count = 0
    for paragraph in doc.paragraphs:
        if not re.match(r"^Tabla\s+\d+\.", paragraph.text.strip()):
            continue
        count += 1
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            set_run_font(run, name="Arial", size=10, italic=True)
    return count


def set_meaningful_alt_text(doc):
    captions = {}
    for index, paragraph in enumerate(doc.paragraphs):
        match = re.match(r"^(Figura\s+\d+\.[^\n]*)", paragraph.text.strip())
        if match:
            captions[index] = match.group(1)

    count = 0
    for index, paragraph in enumerate(doc.paragraphs):
        drawings = paragraph._p.xpath(".//w:drawing")
        if not drawings:
            continue
        candidates = []
        for distance in (1, 2):
            if index - distance in captions:
                candidates.append(captions[index - distance])
            if index + distance in captions:
                candidates.append(captions[index + distance])
        alt = candidates[0] if candidates else "Logotipos institucionales de la portada"
        for doc_pr in paragraph._p.xpath(".//wp:docPr"):
            doc_pr.set("descr", alt)
            doc_pr.set("title", alt.split(".", 1)[0])
            count += 1
    return count


def set_update_fields_on_open(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main():
    doc = Document(SOURCE)

    replacements = [
        (
            "Cup&Cake, es una empresa mexicana fundada en 2014, con más de diez años de operación el crecimiento orgánico le ha permitido subsistir gracias a las compras que actualmente se realizan con base en el empirismo del empresario, ya que lo único que conoce son las fechas de eventos importantes como el Día de San Valentín (14 de febrero), el día del niño (30 de abril), el día de la madre (10 de mayo), entre otros que son los picos de venta que año con año se repiten.",
            "Cup&Cake es una empresa mexicana fundada en 2014. Con más de diez años de operación, su crecimiento orgánico le ha permitido subsistir; sin embargo, las compras se realizan principalmente con base en la experiencia del empresario y en el conocimiento de fechas comerciales recurrentes, como el Día de San Valentín (14 de febrero), el Día del Niño (30 de abril) y el Día de la Madre (10 de mayo), entre otras que generan picos de venta cada año.",
        ),
        (
            "Se aprovecharon los registros históricos de ventas y compras e incorporar variables exógenas disponibles al momento del pronóstico para desarrollar un procedimiento de apoyo a la planeación del presupuesto de abastecimiento.",
            "Se aprovecharon los registros históricos de ventas y compras y se incorporaron variables exógenas disponibles al momento del pronóstico para desarrollar un procedimiento de apoyo a la planeación del presupuesto de abastecimiento.",
        ),
    ]

    prose_changes = replace_text_in_runs(doc, replacements)
    front_matter_changes = fix_static_front_matter(doc)
    removed, retained = reconcile_and_format_references(doc)
    insert_repository_paragraph(doc)
    tables = format_data_tables(doc)
    captions = format_table_captions(doc)
    alt_items = set_meaningful_alt_text(doc)
    set_update_fields_on_open(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)

    print(f"output={OUTPUT}")
    print(f"prose_changes={prose_changes}")
    print(f"front_matter_changes={front_matter_changes}")
    print(f"references_removed={len(removed)}")
    print(f"references_retained={retained}")
    print(f"data_tables_formatted={tables}")
    print(f"captions_formatted={captions}")
    print(f"alt_items_updated={alt_items}")
    for item in removed:
        print(f"REMOVED: {item}")


if __name__ == "__main__":
    main()
