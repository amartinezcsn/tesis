from pathlib import Path
from zipfile import ZipFile
from datetime import datetime
from copy import deepcopy
from lxml import etree as E
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

path=Path('documentacion/TESIS_AGO2026_Rev44_(ZUJ)_12sep2026.docx'); raw=path.read_bytes()
with ZipFile(path) as z: entries=[(i,z.read(i.filename)) for i in z.infolist()]
root=E.fromstring(dict((i.filename,b) for i,b in entries)['word/document.xml'])
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};W='{'+ns['w']+'}'
body=root.find('w:body',ns)
text=lambda p: ''.join(p.xpath('.//w:t/text()',namespaces=ns))
heads=[p for p in body.findall('w:p',ns) if text(p)=='Antecedentes de conjuntos de datos de variables exógenas temporales y comerciales']
assert len(heads)==1
head=heads[0]; position=body.index(head)+1
assert text(body[position]).startswith('Descripción General del Conjunto de Datos de Variables Exógenas Calendáricas')
original=[E.tostring(n) for n in body]
template=next(p for p in body.findall('w:p',ns) if text(p).startswith('El conjunto de ventas de Cup&Cake comenzó'))
def run(s,pr=None):
    r=E.Element(W+'r')
    if pr is not None:r.append(deepcopy(pr))
    t=E.SubElement(r,W+'t');t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=s
    return r
def para(s):
    p=E.Element(W+'p');pr=template.find('w:pPr',ns)
    if pr is not None:p.append(deepcopy(pr))
    p.append(run(s,template.find('w:r/w:rPr',ns)));return p
nodes=[para('Como complemento de los registros internos de ventas y compras, se han reunido archivos con información calendárica, económica y meteorológica relacionada con el contexto del negocio. Estas fuentes se considerarán para examinar si aportan información predictiva al importe semanal de compras, sin presuponer su utilidad ni interpretar las asociaciones encontradas como relaciones causales.'),
para('Su incorporación requerirá verificar la procedencia, la cobertura geográfica y temporal, las unidades y la fecha en que cada dato estuvo disponible. Los eventos de calendario conocidos podrán representar la semana objetivo; los indicadores económicos y meteorológicos observados solo se utilizarán cuando su publicación sea anterior al origen del pronóstico. Los valores proyectados deberán distinguirse de las observaciones históricas.')]
d=Document();t=d.add_table(rows=1,cols=3);t.autofit=False
widths=[3.4,5.4,8.7]
rows=[['Grupo de variables','Información candidata','Condición para su incorporación'],
['Calendario','Festividades y fechas de pago','Documentar la regla de construcción y comprobar que fueran conocidas antes del pronóstico.'],
['Contexto económico','INPC u otros indicadores de precios','Verificar si el archivo contiene niveles del índice o tasas de variación, además de sus fechas de publicación.'],
['Contexto meteorológico','Temperatura histórica o pronósticos archivados','Identificar la cobertura geográfica y utilizar únicamente información disponible en el origen.'],
['Contexto comercial','Eventos o promociones documentados por el negocio','Contar con registros fechados; no deducir retrospectivamente los eventos a partir de las ventas.']]
for i,data in enumerate(rows):
    row=t.rows[0] if i==0 else t.add_row()
    trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    if i==0:trpr.append(OxmlElement('w:tblHeader'))
    for c,s,w in zip(row.cells,data,widths):
        c.text=s;c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        pr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
        for side in ('top','bottom','left','right'):
            e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');m.append(e)
        pr.append(m)
        if i==0:
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'D9E2F3');pr.append(sh)
        for p in c.paragraphs:
            p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=i==0;r.font.color.rgb=RGBColor(0,0,0)
for col,w in zip(t.columns,widths):col.width=Cm(w)
borders=OxmlElement('w:tblBorders')
for side in ('top','bottom','left','right','insideH','insideV'):
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
t._tbl.tblPr.append(borders)
nodes.append(deepcopy(t._tbl))
nodes.append(para('Nota. Las variables descritas son candidatas y no constituyen fuentes aprobadas para el modelado. Su disponibilidad y pertinencia deberán verificarse antes de incorporarlas. La información estatal no se presentará como una medición municipal, ni los valores sintéticos o proyectados como observaciones reales.'))
bookmark='Anexo_Diccionario_Datos'
assert len(root.xpath('//w:bookmarkStart[@w:name="'+bookmark+'"]',namespaces=ns))==1
p=para('El diccionario de datos del ')
f=E.SubElement(p,W+'fldSimple');f.set(W+'instr',' REF '+bookmark+' \\h ');f.set(W+'dirty','true');f.append(run('Anexo A',template.find('w:r/w:rPr',ns)))
p.append(run(' se ampliará con las variables exógenas que se aprueben, documentando su definición, fuente, frecuencia, transformación y disponibilidad temporal. La comparación con y sin estas variables tendrá carácter complementario y no constituirá una segunda hipótesis.',template.find('w:r/w:rPr',ns)))
nodes.append(p)
for i,n in enumerate(nodes):body.insert(position+i,n)
assert original==[E.tostring(n) for n in body if n not in nodes]
out=Path('output/rev44_exogenas');out.mkdir(exist_ok=True)
backup=out/('Rev44_antes_exogenas_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.docx');backup.write_bytes(raw)
candidate=out/'Rev44_exogenas_verificada.docx'
with ZipFile(candidate,'w') as z:
    for i,b in entries:z.writestr(i,E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True) if i.filename=='word/document.xml' else b)
with ZipFile(candidate) as z:
    assert z.testzip() is None
    assert all(z.read(i.filename)==b for i,b in entries if i.filename!='word/document.xml')
assert path.read_bytes()==raw
path.write_bytes(candidate.read_bytes())
print('Aplicados cuatro párrafos y una tabla; referencia REF validada. Respaldo: '+str(backup))
