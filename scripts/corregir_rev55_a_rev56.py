from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.oxml import OxmlElement


SOURCE = Path(r"C:\Python\tesis\documentacion\TESIS_AGO2026_Rev55_(ZUJ).docx")
OUTPUT = Path(r"C:\Python\tesis\documentacion\TESIS_AGO2026_Rev56_(ZUJ).docx")


def normalized(text: str) -> str:
    return " ".join(text.split())


def set_paragraph_text(paragraph, new_text: str) -> None:
    """Replace all inline content while retaining paragraph and first-run formatting."""
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    paragraph.clear()
    run = paragraph.add_run(new_text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def replace_paragraph(document, fragment: str, new_text: str) -> None:
    matches = [p for p in document.paragraphs if fragment in normalized(p.text)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph for {fragment!r}; found {len(matches)}")
    set_paragraph_text(matches[0], new_text)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def delete_reference(document, fragment: str) -> None:
    matches = [p for p in document.paragraphs if fragment in normalized(p.text)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one reference for {fragment!r}; found {len(matches)}")
    delete_paragraph(matches[0])


def insert_reference_before(document, before_fragment: str, text: str) -> None:
    matches = [p for p in document.paragraphs if before_fragment in normalized(p.text)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one insertion point for {before_fragment!r}; found {len(matches)}")
    target = matches[0]
    inserted = target.insert_paragraph_before(text)
    inserted.style = target.style
    inserted.paragraph_format.left_indent = target.paragraph_format.left_indent
    inserted.paragraph_format.first_line_indent = target.paragraph_format.first_line_indent
    inserted.paragraph_format.space_before = target.paragraph_format.space_before
    inserted.paragraph_format.space_after = target.paragraph_format.space_after
    inserted.paragraph_format.line_spacing = target.paragraph_format.line_spacing


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        delete_paragraph(extra)


def find_table_row(document, first_cell_fragment: str):
    matches = []
    for table in document.tables:
        for row in table.rows:
            if first_cell_fragment in normalized(row.cells[0].text):
                matches.append(row)
    if len(matches) != 1:
        raise RuntimeError(f"Expected one table row for {first_cell_fragment!r}; found {len(matches)}")
    return matches[0]


doc = Document(SOURCE)

# Cao et al. (2025): remove the fabricated/mismatched source and recast the claim.
replace_paragraph(
    doc,
    "Una planificación inadecuada de la demanda en el sector de panadería",
    "Una planificación inadecuada de la demanda en el sector de panadería puede generar niveles significativos de merma, con tasas de devolución que oscilan entre el 1.5% y el 19%, lo cual tiene implicaciones económicas y operativas relevantes para las empresas del sector (Hübner et al., 2024). En cuanto a la comparación de métodos, la evidencia de competencias de pronóstico muestra que una mayor complejidad no garantiza una precisión superior y que el desempeño debe evaluarse fuera de muestra para cada contexto (Makridakis et al., 2018; Spiliotis et al., 2025).",
)

# Alekseeva et al. (2021): no matching publication was found. Replace with the
# traceable Latin-American systematic review already present in the bibliography.
replace_paragraph(
    doc,
    "No obstante, gran parte de estas innovaciones han sido implementadas",
    "Las aplicaciones descritas se han estudiado principalmente en contextos con mayor disponibilidad de datos y capacidades tecnológicas. Para las PYMES latinoamericanas, la revisión de Poveda-Valverde y Fierro Barragán (2026) identifica barreras de infraestructura, presupuesto y talento especializado. La evidencia específica sobre microempresas de repostería creativa sigue siendo limitada, por lo que su aplicabilidad debe evaluarse empíricamente y no suponerse a partir de organizaciones de mayor escala.",
)
replace_paragraph(
    doc,
    "Alekseeva, Ginevičius y Stankevičienė (2021) identifican",
    "Poveda-Valverde y Fierro Barragán (2026), mediante una revisión sistemática de aplicaciones de inteligencia artificial en PYMES latinoamericanas, identifican barreras recurrentes relacionadas con infraestructura tecnológica, recursos financieros y disponibilidad de talento especializado.",
)
replace_paragraph(
    doc,
    "A pesar de la creciente disponibilidad de soluciones basadas en Inteligencia Artificial",
    "En el ámbito latinoamericano, la adopción de soluciones basadas en inteligencia artificial por parte de las PYMES enfrenta barreras de infraestructura, presupuesto y talento especializado. Estas condiciones limitan la transferencia directa de aplicaciones desarrolladas en organizaciones con mayores capacidades tecnológicas (Poveda-Valverde & Fierro Barragán, 2026).",
)
replace_paragraph(
    doc,
    "Asimismo, la literatura evidencia que la mayoría de las investigaciones sobre analítica avanzada",
    "Los antecedentes revisados muestran que los resultados obtenidos en organizaciones con capacidades analíticas consolidadas no pueden trasladarse automáticamente a una microempresa. Mikalef et al. (2019) examinan el papel de dichas capacidades en la innovación organizacional, mientras que Poveda-Valverde y Fierro Barragán (2026) documentan barreras de adopción en PYMES latinoamericanas. En conjunto, estas fuentes justifican examinar la aplicabilidad de las herramientas en empresas de menor escala, sin anticipar sus efectos antes de la evaluación empírica.",
)
replace_paragraph(
    doc,
    "Esta investigación propone generar evidencia sobre el uso de herramientas de pronóstico semanal",
    "Esta investigación propone generar evidencia sobre el uso de herramientas de pronóstico semanal en una microempresa con recursos tecnológicos y analíticos limitados. Las restricciones de infraestructura, presupuesto y talento especializado pueden dificultar que las PYMES latinoamericanas adopten soluciones basadas en datos. Por ello, evaluar modelos a partir de la información disponible permitirá identificar alternativas cuya pertinencia pueda valorarse en función de sus capacidades operativas (Poveda-Valverde & Fierro Barragán, 2026).",
)

# Two unsupported Hyndman & Athanasopoulos citations: express the first as a
# bounded result of the literature search and the second as a study decision.
replace_paragraph(
    doc,
    "Dicho vacío se profundiza cuando se examina la integración simultánea",
    "La revisión efectuada no identificó estudios directamente comparables que integren, en una microempresa, el pronóstico semanal del importe de compras y la distribución del presupuesto entre insumos. Esta observación se entiende como resultado de la búsqueda realizada y no como evidencia de ausencia total de investigaciones sobre el tema.",
)
replace_paragraph(
    doc,
    "La ponderación uniforme ofrece igual importancia a las categorías",
    "La ponderación uniforme ofrece igual importancia a las categorías. Si se adoptan otros pesos, deberán justificarse y fijarse antes de evaluar. Además de la medida agregada, revisar errores por insumo permite identificar si una mejora promedio oculta deterioro en categorías concretas. Se informarán por separado las semanas cuyo total cero impida definir una composición observada. Esta regla constituye una decisión de evaluación del estudio.",
)

# Correct the two Maldonado-Guzmán discussions to match traceable publications.
replace_paragraph(
    doc,
    "Por ejemplo, Maldonado Guzmán, Pinzón Castro y García Pérez de Lema (2018)",
    "Valdez-Juárez, García-Pérez-de-Lema y Maldonado-Guzmán (2018) analizaron datos de 412 PYMES industriales y de servicios del noroeste de México. Sus resultados muestran relaciones entre el uso de tecnologías de información y comunicación, la gestión del conocimiento, la innovación y la rentabilidad, por lo que el estudio aporta evidencia nacional sobre capacidades digitales y desempeño empresarial.",
)
replace_paragraph(
    doc,
    "En la misma línea, Maldonado Guzmán y Garza Reyes (2020)",
    "Maldonado-Guzmán y Garza-Reyes (2020) estudiaron la adopción de prácticas de ecoinnovación en la industria automotriz. El trabajo aporta evidencia sobre innovación tecnológica y organizacional en un sector industrial, pero no aborda transformación digital ni modelos predictivos en microempresas.",
)

# Update the state-of-the-art comparison table.
alekseeva_row = find_table_row(doc, "Alekseeva et al. (2021)")
alekseeva_row._tr.getparent().remove(alekseeva_row._tr)

row_2018 = find_table_row(doc, "Maldonado-Guzmán et al. (2018)")
values_2018 = [
    "Valdez-Juárez et al. (2018)",
    "México",
    "TIC; gestión del conocimiento; innovación; rentabilidad",
    "412 PYMES industriales y de servicios del noroeste de México",
    "Encuesta y modelado PLS-SEM",
    "Las TIC se relacionan con la gestión del conocimiento, la innovación y la rentabilidad.",
    "No evalúa IA ni pronósticos.",
    "Aporta contexto nacional sobre capacidades digitales y desempeño empresarial.",
]
for cell, value in zip(row_2018.cells, values_2018):
    set_cell_text(cell, value)

row_2020 = find_table_row(doc, "Maldonado-Guzmán y Garza-Reyes (2020)")
values_2020 = [
    "Maldonado-Guzmán y Garza-Reyes (2020)",
    "México",
    "ecoinnovación; adopción; industria automotriz",
    "Empresas del sector automotriz",
    "Estudio empírico de adopción de prácticas de ecoinnovación",
    "Analiza la adopción de innovaciones tecnológicas y organizacionales en un contexto industrial.",
    "No estudia transformación digital, IA ni pronósticos en microempresas.",
    "Aporta un antecedente mexicano sobre adopción de innovación, con alcance sectorial delimitado.",
]
for cell, value in zip(row_2020.cells, values_2020):
    set_cell_text(cell, value)

# Remove invalid and duplicate bibliography records.
for fragment in [
    "Abrar, M., et al. (2024). An Explainable Multi-Channel",
    "Alekseeva, L., Ginevičius, R., & Stankevičienė, J. (2021). Adoption",
    "Alekseeva, L., Ginevičius, R., & Stankevičienė, J. (2021). Factors",
    "Alekseeva, M., Gorbunova, N., y Parshina, E. (2021)",
    "Cao, G., Duan, Y., Edwards, J. S., & Dwivedi, Y. K. (2025)",
    "Organización para la Cooperación y el Desarrollo Económicos (OCDE). (2022)",
    "PraveenaSri, P., Padma, V., Muruguesan, R., & Usha, S. (2023). Business challenges of forecasting sales in the bakery industry: Applications of machine learning algorithms. ResearchGate Publication.",
]:
    delete_reference(doc, fragment)

# Add the two corrected, traceable Maldonado-Guzmán references in alphabetical order.
insert_reference_before(
    doc,
    "Martins-Turner, K., von Massow",
    "Maldonado-Guzmán, G., & Garza-Reyes, J. A. (2020). Eco-innovation practices’ adoption in the automotive industry. International Journal of Innovation Science, 12(1), 80–98. https://doi.org/10.1108/IJIS-10-2019-0094",
)
insert_reference_before(
    doc,
    "Verhoef, P. C., Broekhuizen",
    "Valdez-Juárez, L. E., García-Pérez-de-Lema, D., & Maldonado-Guzmán, G. (2018). ICT and KM, drivers of innovation and profitability in SMEs. Journal of Information & Knowledge Management, 17(1), 1850007. https://doi.org/10.1142/S0219649218500077",
)

doc.core_properties.title = "Tesis Cup&Cake — Revisión 56"
doc.save(OUTPUT)
print(OUTPUT)
