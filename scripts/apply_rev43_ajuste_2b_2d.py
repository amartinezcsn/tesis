"""Aplica los ajustes 2B, 2C y 2D sobre la Rev42."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from apply_rev42_ajuste_2a import (
    equation_after,
    find_exact,
    find_startswith,
    math_run,
    paragraph_after,
    replace_text,
    subscript,
    subscript_superscript,
)


ROOT = Path(r"C:\Python\tesis")
SOURCE = ROOT / "documentacion" / "TESIS_AGO2026_Rev42_(ZUJ)_05sep2026_DelimitacionCientifica.docx"
OUTPUT = ROOT / "documentacion" / "TESIS_AGO2026_Rev43_(ZUJ)_05sep2026_ProblemaObjetivosHipotesis.docx"
RANKING_FIGURE = ROOT / "output" / "semanal" / "03_ranking_rmse_semanal.png"


def delete_between(start, end) -> None:
    element = start._p.getnext()
    while element is not None and element is not end._p:
        following = element.getnext()
        element.getparent().remove(element)
        element = following


def remove_paragraph(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def upright_subscript(base: str, sub: str):
    return subscript([math_run(base, upright=True)], [math_run(sub, upright=True)])


def forecast_symbol(model_subscript: str):
    accent = OxmlElement("m:acc")
    accent_properties = OxmlElement("m:accPr")
    character = OxmlElement("m:chr")
    character.set(qn("m:val"), "̂")
    accent_properties.append(character)
    accent.append(accent_properties)
    base = OxmlElement("m:e")
    base.append(math_run("y"))
    accent.append(base)
    return subscript([accent], [math_run(model_subscript)])


def error_equation():
    equation = OxmlElement("m:oMath")
    equation.append(subscript([math_run("e")], [math_run("m,t,h")]))
    equation.append(math_run(" = ", upright=True))
    equation.append(subscript([math_run("y")], [math_run("t+h")]))
    equation.append(math_run(" − ", upright=True))
    equation.append(forecast_symbol("m,t+h|t"))
    return equation


def h1_differential_equation():
    equation = OxmlElement("m:oMath")
    equation.append(subscript([math_run("d")], [math_run("m,t,1")]))
    equation.append(math_run(" = ", upright=True))
    equation.append(subscript_superscript([math_run("e")], [math_run("b,t,1")], [math_run("2")]))
    equation.append(math_run(" − ", upright=True))
    equation.append(subscript_superscript([math_run("e")], [math_run("m,t,1")], [math_run("2")]))
    return equation


def h1_hypotheses_equation():
    equation = OxmlElement("m:oMath")
    equation.append(subscript([math_run("H", upright=True)], [math_run("0,m")]))
    equation.append(math_run(":  E", upright=True))
    equation.append(math_run("[", upright=True))
    equation.append(subscript([math_run("d")], [math_run("m,t,1")]))
    equation.append(math_run("] ≤ 0;    ", upright=True))
    equation.append(subscript([math_run("H", upright=True)], [math_run("1,m")]))
    equation.append(math_run(":  E", upright=True))
    equation.append(math_run("[", upright=True))
    equation.append(subscript([math_run("d")], [math_run("m,t,1")]))
    equation.append(math_run("] > 0", upright=True))
    return equation


def h2_differential_equation():
    equation = OxmlElement("m:oMath")
    equation.append(
        subscript_superscript(
            [math_run("d")],
            [math_run("m,t,h")],
            [math_run("exo", upright=True)],
        )
    )
    equation.append(math_run(" = ", upright=True))
    equation.append(
        subscript_superscript(
            [math_run("e")],
            [math_run("m,t,h,hist", upright=True)],
            [math_run("2")],
        )
    )
    equation.append(math_run(" − ", upright=True))
    equation.append(
        subscript_superscript(
            [math_run("e")],
            [math_run("m,t,h,exo", upright=True)],
            [math_run("2")],
        )
    )
    return equation


def h2_hypotheses_equation():
    equation = OxmlElement("m:oMath")
    equation.append(subscript_superscript([math_run("H", upright=True)], [math_run("0")], [math_run("exo", upright=True)]))
    equation.append(math_run(":  E", upright=True))
    equation.append(math_run("[", upright=True))
    equation.append(subscript_superscript([math_run("d")], [math_run("m,t,h")], [math_run("exo", upright=True)]))
    equation.append(math_run("] ≤ 0;    ", upright=True))
    equation.append(subscript_superscript([math_run("H", upright=True)], [math_run("2")], [math_run("exo", upright=True)]))
    equation.append(math_run(":  E", upright=True))
    equation.append(math_run("[", upright=True))
    equation.append(subscript_superscript([math_run("d")], [math_run("m,t,h")], [math_run("exo", upright=True)]))
    equation.append(math_run("] > 0", upright=True))
    return equation


def selection_equation():
    equation = OxmlElement("m:oMath")
    equation.append(subscript_superscript([math_run("m")], [math_run("h")], [math_run("*")]))
    equation.append(math_run(" = ", upright=True))
    equation.append(subscript([math_run("arg min", upright=True)], [math_run("m")]))
    equation.append(math_run(" ", upright=True))
    equation.append(subscript([math_run("RMSE", upright=True)], [math_run("m,h")]))
    return equation


def adoption_equation():
    equation = OxmlElement("m:oMath")
    selected_model = subscript_superscript([math_run("m")], [math_run("h")], [math_run("*")])
    equation.append(subscript([math_run("RMSE", upright=True)], [selected_model, math_run(",h")]))
    equation.append(math_run(" < ", upright=True))
    equation.append(subscript([math_run("RMSE", upright=True)], [math_run("b,h")]))
    equation.append(math_run("    ∧    ", upright=True))
    equation.append(subscript([math_run("p", upright=True)], [math_run("Holm", upright=True)]))
    equation.append(math_run(" ≤ 0.05", upright=True))
    return equation


def add_numbered_objective(document: Document, anchor, number: int, text: str):
    paragraph = paragraph_after(document, anchor, f"{number}. {text}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.first_line_indent = Inches(-0.3)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def replace_figure_before_caption(caption, image_path: Path, width: float = 6.2) -> None:
    image_paragraph_element = caption._p.getprevious()
    if image_paragraph_element is None:
        raise ValueError(f"No existe una imagen antes de {caption.text}")
    from docx.text.paragraph import Paragraph

    paragraph = Paragraph(image_paragraph_element, caption._parent)
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))


def update_results_table(document: Document) -> None:
    table = document.tables[26]
    values = [
        ["Horizonte", "Modelo y representación", "MAE", "RMSE", "Interpretación"],
        ["h=1", "Ingenuo estacional de 52 semanas", "438.19", "695.19", "Menor RMSE descriptivo"],
        ["h=4", "Hurdle HistGradientBoosting con exógenas", "302.64", "382.81", "Menor RMSE descriptivo"],
    ]
    while len(table.rows) > len(values):
        table._tbl.remove(table.rows[-1]._tr)
    for row, row_values in zip(table.rows, values):
        for cell, value in zip(row.cells, row_values):
            cell.text = value
            cell.vertical_alignment = 1
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in table.rows[0].cells[0].paragraphs[0].runs:
        run.bold = True
    for cell in table.rows[0].cells[1:]:
        for run in cell.paragraphs[0].runs:
            run.bold = True


def main() -> None:
    document = Document(SOURCE)

    # 2B: planteamiento del problema y pregunta de investigación.
    problem = find_exact(document, "El planteamiento del problema")
    context = find_exact(document, "Contexto del caso de estudio")
    delete_between(problem, context)
    current = problem
    for text in (
        "Las microempresas suelen planear sus compras con recursos analíticos limitados y una fuerte dependencia de la experiencia del propietario. Las restricciones de infraestructura, presupuesto y capital humano dificultan la adopción de soluciones avanzadas de inteligencia artificial y reducen la capacidad de convertir registros operativos en pronósticos reproducibles (Alekseeva et al., 2021; Mikalef et al., 2019).",
        "En las microempresas alimentarias, la decisión inmediata no consiste únicamente en anticipar la demanda, sino en reservar recursos monetarios para adquirir insumos bajo variabilidad, estacionalidad y periodos sin movimiento. Cuando el importe futuro de compras se estima sólo mediante memoria o intuición, no existe una medida verificable del error ni una base común para comparar alternativas de planeación.",
        "La mayor complejidad algorítmica no garantiza una mayor precisión. En series cortas e intermitentes, los modelos estadísticos y de aprendizaje automático pueden sobreajustarse o aprender discontinuidades producidas por ausencia de registro. Por ello, su utilidad debe demostrarse frente a reglas empíricas reproducibles y sobre semanas que no participaron en el entrenamiento (Makridakis et al., 2018; Spiliotis et al., 2022).",
        "La literatura revisada ofrece evidencia amplia en cadenas minoristas y colecciones extensas de productos, pero evidencia limitada sobre comparaciones temporales reproducibles para pronosticar el importe semanal de compras en una sola microempresa de repostería con Small Data. El vacío se concentra en determinar si la complejidad estadística o algorítmica aporta una reducción verificable del error en este contexto específico.",
        "En consecuencia, el problema científico es evaluar, bajo el mismo protocolo temporal, cuánto error fuera de muestra presentan las líneas base, los modelos estadísticos y los algoritmos de aprendizaje automático al estimar el importe semanal de compras. La respuesta permitirá establecer una regla de selección para el presupuesto de abastecimiento sin atribuir efectos no medidos sobre inventarios, merma, rentabilidad o desabasto.",
    ):
        current = paragraph_after(document, current, text)
        current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    delimiter = find_exact(document, "Delimitación científica y unidad de análisis")
    delete_between(context, delimiter)
    current = context
    for text in (
        "Cup&Cake es una microempresa de repostería creativa fundada en 2014 en Tizayuca, Hidalgo. Su operación combina productos personalizados, pedidos asociados con celebraciones y fechas comerciales y compras de insumos realizadas conforme a las necesidades observadas del negocio.",
        "La planeación del abastecimiento se ha apoyado principalmente en la experiencia acumulada del propietario, el conocimiento de eventos recurrentes y la revisión de movimientos recientes. Este conocimiento es valioso para interpretar la operación, pero no produce por sí mismo un error cuantificable ni permite saber si otra técnica generaliza mejor.",
        "La empresa dispone de registros diarios de ventas y compras que permiten construir una serie semanal y evaluar alternativas de pronóstico. Su necesidad operativa es estimar el importe probable de compras de las siguientes una a cuatro semanas para revisar y reservar presupuesto de corto plazo.",
    ):
        current = paragraph_after(document, current, text)
        current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    question = find_startswith(document, "¿En qué medida los modelos estadísticos")
    replace_text(
        question,
        "¿En qué medida las configuraciones estadísticas y de aprendizaje automático reducen el error de "
        "pronóstico fuera de muestra respecto de la línea base de último valor semanal observado al estimar "
        "el importe semanal de compras de Cup&Cake, considerando h=1 como horizonte principal y h=4 como "
        "horizonte complementario?",
    )
    question.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 2C: objetivo general y objetivos específicos.
    general = find_startswith(document, "Evaluar el desempeño de modelos estadísticos")
    replace_text(
        general,
        "Evaluar si los modelos estadísticos y de aprendizaje automático mejoran la precisión del pronóstico "
        "fuera de muestra del importe semanal de compras frente a la línea base de último valor observado, "
        "mediante validación temporal de ventana deslizante, para apoyar la planeación del presupuesto de "
        "abastecimiento de Cup&Cake.",
    )
    general.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    specifics = find_exact(document, "Objetivos específicos")
    justification = find_exact(document, "Justificación")
    delete_between(specifics, justification)
    current = specifics
    objectives = (
        "Construir y auditar una serie semanal del importe real de compras a partir de los registros diarios disponibles.",
        "Caracterizar la cobertura, tendencia, estacionariedad, dependencia temporal, intermitencia, dispersión y valores atípicos de la serie objetivo.",
        "Construir variables históricas, calendáricas y exógenas respetando su disponibilidad temporal y evitando fuga de información.",
        "Entrenar y ajustar líneas base, modelos estadísticos y algoritmos de aprendizaje automático para los horizontes directos h=1, h=2, h=3 y h=4 mediante ventanas de entrenamiento de 52 semanas.",
        "Comparar la precisión y estabilidad fuera de muestra mediante RMSE, MAE, métricas escaladas y contrastes estadísticos, con h=1 como horizonte principal y h=4 como evidencia complementaria.",
        "Integrar predicciones, métricas, contrastes y criterios de selección en un tablero que apoye el presupuesto semanal de abastecimiento y su consolidación de cuatro semanas.",
    )
    for number, text in enumerate(objectives, 1):
        current = add_numbered_objective(document, current, number, text)

    # 2D: hipótesis, formulación matemática y regla de decisión.
    hypotheses = find_exact(document, "Hipótesis")
    theoretical = find_exact(document, "MARCO TEÓRICO")
    delete_between(hypotheses, theoretical)

    current = paragraph_after(
        document,
        hypotheses,
        "Sea ŷₘ,ₜ₊ₕ|ₜ el pronóstico del modelo m emitido en el origen t para la semana t+h. "
        "El error fuera de muestra se define como:",
    )
    current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    current = equation_after(document, current, error_equation())

    h1_heading = paragraph_after(document, current, "Hipótesis de investigación H1", style="Heading 4")
    h1_heading.paragraph_format.keep_with_next = True
    current = paragraph_after(
        document,
        h1_heading,
        "H1. Al menos una configuración estadística o de aprendizaje automático presentará un error "
        "cuadrático esperado fuera de muestra significativamente menor que la línea base primaria de último "
        "valor semanal observado para el horizonte principal de una semana. Si b identifica la línea base, "
        "el diferencial de pérdida se define como:",
    )
    current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    current = equation_after(document, current, h1_differential_equation())
    current = paragraph_after(
        document,
        current,
        "Un diferencial positivo favorece al modelo evaluado. Para cada configuración se contrastarán las "
        "siguientes hipótesis unilaterales:",
    )
    current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    current = equation_after(document, current, h1_hypotheses_equation())
    current = paragraph_after(
        document,
        current,
        "La inferencia utilizará la prueba Diebold-Mariano unilateral con corrección de muestra finita. "
        "Debido a las comparaciones múltiples frente a una misma referencia, los valores p se ajustarán con "
        "el procedimiento Holm y H1 sólo se considerará respaldada cuando al menos una diferencia favorable "
        "mantenga p ajustado no mayor que 0.05.",
    )
    current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    h2_heading = paragraph_after(document, current, "Hipótesis de investigación H2", style="Heading 4")
    h2_heading.paragraph_format.keep_with_next = True
    current = paragraph_after(
        document,
        h2_heading,
        "H2. La incorporación de variables exógenas disponibles antes de la semana pronosticada reducirá el "
        "error cuadrático esperado fuera de muestra respecto de configuraciones equivalentes construidas "
        "únicamente con información histórica. El diferencial se define como:",
    )
    current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    current = equation_after(document, current, h2_differential_equation())
    current = paragraph_after(
        document,
        current,
        "Para una misma familia de modelo y horizonte se examinarán las siguientes hipótesis:",
    )
    current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    current = equation_after(document, current, h2_hypotheses_equation())
    current = paragraph_after(
        document,
        current,
        "Los contrastes de H2 se interpretarán como evidencia exploratoria por modelo y horizonte. Sus valores "
        "p no se utilizarán como confirmación general mientras no se aplique una corrección explícita por "
        "multiplicidad.",
    )
    current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    decision_heading = paragraph_after(document, current, "Regla de decisión predictiva", style="Heading 4")
    decision_heading.paragraph_format.keep_with_next = True
    current = paragraph_after(
        document,
        decision_heading,
        "Para cada horizonte, la configuración descriptivamente preferida será la que minimice el RMSE:",
    )
    current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    current = equation_after(document, current, selection_equation())
    current = paragraph_after(
        document,
        current,
        "Una configuración avanzada sólo se recomendará como sustituta de la línea base primaria cuando "
        "cumpla simultáneamente las dos condiciones siguientes:",
    )
    current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    current = equation_after(document, current, adoption_equation())
    current = paragraph_after(
        document,
        current,
        "Si las condiciones no se cumplen, el tablero conservará la línea base como referencia operativa y "
        "presentará cualquier mejora únicamente como evidencia descriptiva. Esta regla permite obtener una "
        "recomendación útil incluso cuando un modelo complejo no demuestra superioridad estadística.",
    )
    current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Consistencia de la formulación central y de los resultados regenerados.
    replace_text(
        find_startswith(document, "El análisis identificó una serie errática"),
        "El análisis identificó una serie errática, con ADI de 1.16, CV² de 1.03 y 13.5% de semanas con "
        "importe cero dentro del bloque de desarrollo. En h=1, el ingenuo estacional de 52 semanas obtuvo el "
        "menor RMSE descriptivo (695.19) y MAE de 438.19; en h=4, la mejor configuración descriptiva fue "
        "Hurdle HistGradientBoosting con variables exógenas, con RMSE de 382.81 y MAE de 302.64. Ninguna "
        "configuración estadística o de aprendizaje automático superó significativamente la línea base "
        "primaria después del ajuste Holm. Los resultados se comunican mediante un sistema de soporte a la "
        "decisión que conserva la interpretación humana.",
    )
    replace_text(
        find_startswith(document, "This research evaluates statistical"),
        "This research evaluates statistical and machine-learning models for weekly procurement-budget "
        "forecasting at Cup&Cake, a small bakery in Tizayuca, Hidalgo. Daily sales and purchasing records "
        "were aggregated into a weekly series and complemented with variables available before each forecast. "
        "The final 16 temporal origins were reserved for evaluation. At h=1, the 52-week seasonal naive "
        "benchmark achieved the lowest descriptive RMSE (695.19) and an MAE of 438.19; at h=4, exogenous "
        "Hurdle HistGradientBoosting achieved an RMSE of 382.81 and an MAE of 302.64. No statistical or "
        "machine-learning configuration significantly outperformed the primary baseline after Holm "
        "adjustment. Results support human review through an auditable decision-support dashboard.",
    )
    replace_text(
        find_startswith(document, "Si la introducción de algoritmos avanzados"),
        "Una reducción descriptiva del error no demuestra por sí misma una mejora operativa o financiera. "
        "En esta investigación, la superioridad predictiva requiere una diferencia favorable de error fuera "
        "de muestra y significancia unilateral después del ajuste Holm. Incluso cuando se cumpla este criterio, "
        "el resultado se interpretará como evidencia sobre precisión del pronóstico y no como prueba causal "
        "de reducción de merma, incremento de rentabilidad o mejora del flujo de efectivo.",
    )
    replace_text(
        find_startswith(document, "El DSS debe apoyar la revisión del presupuesto"),
        "El DSS debe apoyar la revisión del presupuesto y mostrar la incertidumbre asociada a cada horizonte. "
        "En h=1 comunica al ingenuo estacional de 52 semanas como referencia con menor RMSE descriptivo; en "
        "h=4, Hurdle HistGradientBoosting con variables exógenas obtuvo RMSE de 382.81 y MAE de 302.64. "
        "Ninguno de estos resultados autoriza automatizar compras ni afirmar superioridad estadística frente "
        "a la referencia primaria.",
    )
    replace_text(
        find_startswith(document, "La evaluación comparativa se reporta por horizonte"),
        "La evaluación comparativa se reporta por horizonte. En h=1, el ingenuo estacional de 52 semanas "
        "registró RMSE de 695.19 y MAE de 438.19, frente a RMSE de 1029.47 y MAE de 758.65 de la línea base "
        "primaria de último valor observado. En h=4, Hurdle HistGradientBoosting con variables exógenas obtuvo "
        "RMSE de 382.81 y MAE de 302.64, frente a RMSE de 866.60 y MAE de 601.52 de la línea base primaria. "
        "Las métricas no se combinan entre horizontes.",
    )
    replace_text(
        find_startswith(document, "H1 se contrastó principalmente"),
        "H1 se contrastó principalmente en h=1 contra la línea base de último valor semanal observado. "
        "Aunque varias configuraciones redujeron descriptivamente el RMSE, ninguna mantuvo una diferencia "
        "significativa después de la prueba Diebold-Mariano unilateral y el ajuste Holm; en h=4 tampoco se "
        "obtuvieron contrastes significativos. Por ello, H1 no queda respaldada. En H2 se observaron diferencias "
        "unilaterales favorables para Ridge y Ridge con transformación log1p en comparaciones específicas, pero "
        "se interpretan como evidencia exploratoria sin corrección por multiplicidad.",
    )
    replace_text(
        find_startswith(document, "Los resultados se reportan exclusivamente"),
        "Los resultados se reportan exclusivamente para el importe semanal de compras. En h=1, el ingenuo "
        "estacional de 52 semanas obtuvo el menor RMSE descriptivo; en h=4, Hurdle HistGradientBoosting con "
        "variables exógenas obtuvo el menor RMSE descriptivo. MAPE se mantiene como indicador secundario debido "
        "a la presencia de semanas con importe cero.",
    )
    replace_text(
        find_startswith(document, "Los hallazgos no respaldan una superioridad general"),
        "Los hallazgos no respaldan la superioridad inferencial de un modelo complejo. En h=1, una referencia "
        "estacional simple obtuvo el menor RMSE descriptivo; en h=4, una configuración Hurdle "
        "HistGradientBoosting con exógenas obtuvo el menor error, pero no superó significativamente la línea "
        "base primaria tras el ajuste Holm. Esta diferencia entre horizontes impide trasladar conclusiones de "
        "uno a otro.",
    )
    replace_text(
        find_startswith(document, "La evidencia no respalda H1"),
        "La evidencia no respalda H1 bajo el criterio inferencial predefinido. En h=1, el ingenuo estacional de "
        "52 semanas obtuvo el menor RMSE descriptivo; en h=4, Hurdle HistGradientBoosting con variables "
        "exógenas produjo el menor error descriptivo. Sin embargo, ninguna configuración estadística o de "
        "aprendizaje automático superó la prueba Diebold-Mariano con ajuste Holm frente a la línea base "
        "primaria, por lo que no se afirma superioridad general.",
    )

    # Evidencia visual y tabla final alineadas con el experimento semanal.
    figure_31 = find_exact(document, "Figura 31 - Mejora relativa del RMSE frente a la línea base empírica.")
    replace_figure_before_caption(figure_31, RANKING_FIGURE)
    replace_text(figure_31, "Figura 31 - Ranking de RMSE por horizonte y configuración")
    figure_32 = find_exact(document, "Figura 32 - Mejor modelo por objetivo y representación de datos.")
    image_32 = figure_32._p.getprevious()
    if image_32 is not None:
        image_32.getparent().remove(image_32)
    remove_paragraph(figure_32)

    for paragraph in list(document.paragraphs):
        if paragraph.style.name == "table of figures" and paragraph.text.startswith("Figura 31 -"):
            replace_text(paragraph, "Figura 31 - Ranking de RMSE por horizonte y configuración")
        elif paragraph.style.name == "table of figures" and paragraph.text.startswith("Figura 32 -"):
            remove_paragraph(paragraph)
        elif paragraph.style.name == "table of figures" and paragraph.text.startswith("Tabla 11 -"):
            replace_text(paragraph, "Tabla 11 - Resultados principales por horizonte")

    table_caption = find_exact(document, "Tabla 11 - Modelos ganadores y desempeño temporal")
    replace_text(table_caption, "Tabla 11 - Resultados principales por horizonte")
    update_results_table(document)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
