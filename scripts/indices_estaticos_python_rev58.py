from pathlib import Path
import re
import unicodedata

import pypdfium2 as pdfium
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

SOURCE_DOCX = Path(r"C:\Python\tesis\documentacion\TESIS_SEP2026_Rev58_(ZUJ)_Corregida.docx")
DOCX = Path(r"C:\Python\tesis\documentacion\TESIS_SEP2026_Rev58_(ZUJ)_Final.docx")
PDF = Path(r"C:\Python\tesis\tmp\rev58_render\TESIS_SEP2026_Rev58_QA.pdf")


def norm(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.replace("�", "")
    return re.sub(r"\s+", " ", value).strip().casefold()


def insert_after(paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new.style = style
        except KeyError:
            pass
    new.add_run(text)
    return new


doc = Document(SOURCE_DOCX)
pdf = pdfium.PdfDocument(str(PDF))
page_texts = [norm(pdf[i].get_textpage().get_text_range()) for i in range(len(pdf))]

captions = []
for p in doc.paragraphs:
    text = p.text.strip()
    if not re.match(r"^(Figura|Tabla|Imagen)\s+\d+\s*[.\-]", text, re.I):
        continue
    if text.lower().startswith("fuente:"):
        continue
    needle = norm(text)[:55]
    page = None
    for idx, page_text in enumerate(page_texts, start=1):
        if needle and needle in page_text:
            page = max(1, idx - 10)  # main matter restarts at 1 on physical PDF page 11
            break
    if page is None:
        raise RuntimeError(f"Caption page not found: {text}")
    captions.append((text, page))

paras = doc.paragraphs
index_heading = next(p for p in paras if p.text.strip() == "ÍNDICE")
figure_heading = next(p for p in paras if p.text.strip() == "TABLAS, IMÁGENES Y ECUACIONES")
summary = next(p for p in paras if p.text.strip() == "RESUMEN")

# Remove the legacy empty table-of-figures field/result.
node = figure_heading._p.getnext()
while node is not None and node is not summary._p:
    following = node.getnext()
    node.getparent().remove(node)
    node = following

figure_heading.text = "ÍNDICE DE TABLAS Y FIGURAS"
anchor = figure_heading
for text, page in captions:
    anchor = insert_after(anchor, f"{text}\t{page}", "table of figures")

summary.paragraph_format.page_break_before = True

# The existing main TOC is now static and correct except for the legacy error prefix.
for p in doc.paragraphs:
    if p.text.startswith("No se encontraron entradas de tabla de contenido."):
        p.text = p.text.replace("No se encontraron entradas de tabla de contenido.", "", 1)
        break

# The inserted table/figure index occupies two front-matter pages; the main matter keeps its restart.
for p in doc.paragraphs:
    if p.style and p.style.name.lower() == "toc 2" and p.text.startswith("RESUMEN\t"):
        p.text = "RESUMEN\t10"
        break

doc.save(DOCX)
print(f"saved={DOCX} captions={len(captions)}")
for text, page in captions:
    print(f"{page}\t{text}")
