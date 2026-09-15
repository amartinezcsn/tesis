from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy
from datetime import datetime
import hashlib, json
from lxml import etree as E
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from openpyxl import load_workbook

path=Path('documentacion/TESIS_AGO2026_Rev44_(ZUJ)_12sep2026.docx')
raw=path.read_bytes(); out=Path('output/rev44_diccionario_20260913'); out.mkdir(exist_ok=True)
with ZipFile(path) as z: entries=[(i,z.read(i.filename)) for i in z.infolist()]
root=E.fromstring(dict((i.filename,b) for i,b in entries)['word/document.xml'])
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}; W='{'+ns['w']+'}'
body=root.find('w:body',ns); ps=body.findall('w:p',ns)
def txt(p): return ''.join(p.xpath('.//w:t/text()',namespaces=ns))
def find(s):
    found=[p for p in ps if txt(p).strip()==s]; assert len(found)==1,(s,len(found)); return found[0]
assert not any(txt(p)=='Anexo A Diccionario de datos' for p in ps)
saleshead=find('Descripción General del Conjunto de Datos de Ventas')
buyhead=find('Descripción General del Conjunto de Datos de Compras (Egresos)')
endhead=find('Antecedentes de conjuntos de datos de variables exógenas temporales y comerciales')
normal=find('El conjunto de ventas tiene su origen en un archivo de Microsoft Excel. Posteriormente, se incorporaron los reportes descargados de KYTE, para lo cual se ajustó la estructura del archivo existente. Excel se mantuvo como medio de consolidación de ambas procedencias. Antes del análisis, se revisará la correspondencia de los campos y se comprobará que la integración no haya generado duplicidades o inconsistencias.')
before_eq=[E.tostring(n) for n in root.xpath('//m:oMath',namespaces=ns)]
ref_start=list(body).index(find('REFERENCIAS'))
refs=[E.tostring(n) for n in list(body)[ref_start:] if n.tag!=W+'sectPr']
met=list(body).index(find('METODOLOGÍA')); dev=list(body).index(find('DESARROLLO'))
method=[E.tostring(n) for n in list(body)[met:dev]]
changes=[]
def settext(p,s):
    changes.append({'before':txt(p),'after':s})
    ts=p.xpath('.//w:t',namespaces=ns); assert ts
    ts[0].text=s
    for t in ts[1:]:t.text=''
def para(s,template=None):
    p=E.Element(W+'p'); t=template if template is not None else normal
    pr=t.find('w:pPr',ns)
    if pr is not None:p.append(deepcopy(pr))
    r=E.SubElement(p,W+'r'); rp=t.find('w:r/w:rPr',ns)
    if rp is not None:r.append(deepcopy(rp))
    E.SubElement(r,W+'t').text=s
    return p
def replace_between(a,b,strings):
    start=list(body).index(a)+1; end=list(body).index(b)
    removed=list(body)[start:end]
    assert not any(n.xpath('.//m:oMath|.//w:drawing|.//w:sectPr',namespaces=ns) for n in removed)
    for n in removed:body.remove(n)
    for i,s in enumerate(strings):body.insert(start+i,para(s))

for p in ps:
    s=txt(p)
    if s.startswith('En 2022, Cup&Cake inició el registro de sus ventas'):
        settext(p,s.replace('En 2022,','En noviembre de 2020,',1))
    elif s.startswith('Aunque Cup&Cake inició operaciones en 2012, el registro'):
        settext(p,s.replace('comenzó en 2022','comenzó en noviembre de 2020'))
    elif s.startswith('Para registrar las compras de insumos y materias primas, se desarrolló'):
        settext(p,s.replace('Para registrar','En diciembre de 2016, para registrar',1))
    elif s.startswith('Por su parte, el registro de compras surgió'):
        settext(p,'En diciembre de 2016, Cup&Cake comenzó a registrar las compras de insumos y materias primas mediante un sistema propio desarrollado en Microsoft Access. Esta herramienta permite conservar, consultar y extraer la información de las adquisiciones para su revisión y análisis.')
    elif s.strip()=='El registro de ventas y compra de insumos':
        settext(p,'El registro de ventas y compras de insumos')
    elif s.startswith('Nota. Elaboración propia con base en la información proporcionada sobre Cup&Cake.'):
        settext(p,'Nota. Elaboración propia con base en la información proporcionada sobre Cup&Cake. El inicio de los registros no garantiza cobertura continua. Los valores iguales a cero no necesariamente indican ausencia de ventas o compras; también podrían corresponder a omisiones o información incompleta. Su significado se verificará antes del análisis.')

replace_between(saleshead,buyhead,[
'El conjunto de ventas de Cup&Cake comenzó a registrarse en noviembre de 2020 mediante Microsoft Excel. Posteriormente, se incorporaron los reportes descargados de KYTE al archivo existente, cuya estructura se adaptó para consolidar la información. Las ventas podrán utilizarse como predictores complementarios, siempre que estén disponibles antes de cada pronóstico.',
'El archivo Ventas.xlsx contiene 25 campos. Su definición, tipo esperado y reglas de revisión se presentan en el Anexo A. Antes de agregar por semana se verificarán los importes, las fechas, los estados y la identificación de las transacciones para evitar duplicidades.'
])
replace_between(buyhead,endhead,[
'El conjunto de compras procede de un sistema propio desarrollado en Microsoft Access, utilizado desde diciembre de 2016 para registrar adquisiciones de insumos y materias primas. Constituirá la fuente principal para calcular el importe total semanal de compras y su distribución porcentual entre los principales insumos.',
'El archivo Compras.xlsx contiene 10 campos. El Anexo A documenta sus nombres reales y su correspondencia con las variables del pipeline. Se revisarán los importes y las fechas, se homologarán las descripciones y se delimitarán las adquisiciones incluidas. La existencia de registros no garantiza cobertura continua ni permite interpretar automáticamente los valores cero como ausencia de compras.'
])
for p in ps:
    s=txt(p)
    if s.startswith('La Tabla 2 resumirá la procedencia'):
        settext(p,'La Tabla 3 resume la procedencia y la función de ambos conjuntos. El Anexo A presenta el diccionario de campos de origen y variables analíticas. El periodo utilizable se establecerá después de verificar la cobertura, sin equiparar las fechas de inicio del registro con una muestra completa.')
    elif s.startswith('Como se observa en la Tabla 3 - Conjuntos de datos'):
        settext(p,'Las compras definen las variables objetivo y las ventas constituyen una fuente predictiva complementaria. La agregación semanal no supone una correspondencia directa entre cada venta y cada adquisición ni permite, por sí sola, evaluar rentabilidad o existencias de inventario.')

def table(headers,rows,widths):
    d=Document(); t=d.add_table(rows=1,cols=len(headers)); t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for i,s in enumerate(headers):t.rows[0].cells[i].text=s
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=str(s)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for c,w in zip(row.cells,widths):
            c.width=Cm(w); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
            for name in ('top','left','bottom','right'):
                n=OxmlElement('w:'+name);n.set(qn('w:w'),'90');n.set(qn('w:type'),'dxa');margins.append(n)
            pr.append(margins)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'D9E2F3');pr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=(ri==0);r.font.color.rgb=RGBColor(0,0,0)
    borders=OxmlElement('w:tblBorders')
    for name in ('top','left','bottom','right','insideH','insideV'):
        n=OxmlElement('w:'+name);n.set(qn('w:val'),'single');n.set(qn('w:sz'),'4');n.set(qn('w:color'),'D9D9D9');borders.append(n)
    t._tbl.tblPr.append(borders)
    return deepcopy(t._tbl)

oldtables=[n for n in body.findall('w:tbl',ns) if 'Number / NumPedido / Id' in txt(n)]
assert len(oldtables)==1
old=oldtables[0]; loc=list(body).index(old)
body.remove(old);body.insert(loc,table(['Aspecto','Ventas','Compras'],[
('Inicio del registro','Noviembre de 2020','Diciembre de 2016'),
('Procedencia','Excel y reportes descargados de KYTE','Sistema propio en Microsoft Access'),
('Archivo inspeccionado','Ventas.xlsx, hoja Hoja1, 25 campos','Compras.xlsx, hoja Hoja1, 10 campos'),
('Información principal','Identificación, fechas, importes y productos vendidos','Fechas, importes, cantidades y descripciones de insumos'),
('Preparación requerida','Compatibilidad entre fuentes, estados, importes y duplicidades','Fechas, importes, catálogo y delimitación de las compras'),
('Función en el estudio','Predictores complementarios disponibles antes del pronóstico','Importe total semanal y participaciones monetarias por insumo'),
('Frecuencia analítica','Semanal','Semanal')],[3.3,7.1,7.1]))
body.insert(loc+1,para('Nota. Elaboración propia a partir de los archivos inspeccionados y la información histórica de Cup&Cake. Los ceros requieren verificación; no equivalen automáticamente a ausencia de operaciones. El diccionario detallado se presenta en el Anexo A.'))

# One row per physical source field; no fabricated aliases or sample personal data.
sales=[
('Number','Texto','Identificador registrado de la transacción.','Control; comprobar unicidad, no usar como predictor.'),
('Status','Categoría','Estado registrado de la venta.','Validar categorías y criterio de inclusión; no asumir que todos los estados son ventas concluidas.'),
('Fecha/Hora','Fecha y hora','Marca temporal del registro.','Convertir y contrastar con Fecha; no equivale necesariamente a fecha de disponibilidad.'),
('Cantidad','Número','Cantidad consignada en la venta.','Unidad y diferencia respecto de Total de ítems pendientes de verificación.'),
('Total de ítems','Número','Total de ítems consignado.','Verificar si cuenta unidades o líneas del pedido; no asumir equivalencia con Cantidad.'),
('Descri. Items','Texto','Descripción de los productos de la transacción.','Descripción auxiliar; no se modela texto libre en la versión inicial.'),
('Subtotal','Número monetario','Subtotal registrado antes de ajustes.','MXN por confirmar en origen; revisar conciliación, no sumar al Importe.'),
('Descuento','Número','Descuento registrado.','Confirmar si expresa MXN o porcentaje antes de operar con él.'),
('Tasa','Número','Campo de tasa del reporte.','Definición y unidad pendientes; no asumir importe de impuesto.'),
('Envío','Número monetario','Cargo de envío registrado.','Confirmar inclusión en Importe; no duplicar cargos.'),
('Importe','Número monetario','Importe registrado de la venta.','Candidato para agregado nominal semanal; validar concepto, moneda, ceros y estados.'),
('Ganancia','Número','Campo de ganancia registrado.','Cálculo pendiente de verificación; no equivale a utilidad neta demostrada ni es objetivo del estudio.'),
('Nombre','Texto','Nombre asociado al registro.','Rol pendiente de confirmar; dato potencialmente personal, excluido de predictores.'),
('Vendedor','Texto o categoría','Responsable o canal consignado.','Dato administrativo; excluir identificadores personales del modelado.'),
('Observación','Texto','Notas de la transacción.','Puede contener datos personales; no publicar valores ni usar como predictor inicial.'),
('Fecha','Fecha','Fecha consignada sin hora.','Normalizar y contrastar con Fecha/Hora; definir fecha válida para agregación.'),
('Año','Entero','Año consignado.','Contrastar con la fecha validada; preferir derivación consistente.'),
('Mes','Entero','Mes consignado.','Valores esperados 1 a 12; contrastar con la fecha.'),
('Día','Entero','Día del mes consignado.','Validar según mes y año.'),
('Pastel','Indicador','Indicador de categoría pastel.','Se observaron 0, 1 y texto null; verificar codificación. No convertir null en 0.'),
('Galletas','Indicador','Indicador de categoría galletas.','Se observaron 0, 1 y texto null; no prueba cantidades vendidas.'),
('Otros','Indicador','Indicador de otros productos de venta.','No confundir con otros de la composición de compras; revisar null.'),
('Cupcakes','Indicador','Indicador de categoría cupcakes.','Se observaron 0, 1 y texto null; verificar codificación.'),
('Pedido','Texto','Referencia interna del pedido.','Contiene tipos mixtos; conservar como identificador y verificar relación con Number.'),
('NumPedido','Texto','Número o referencia del pedido.','Comprobar relación con Pedido y Number; la unicidad no está garantizada.')]
buys=[
('PROVEEDOR','Texto','Proveedor consignado en la compra.','Control administrativo; no es predictor inicial.'),
('FECHA','Fecha','Fecha registrada de adquisición.','Se homologa a fecha; fecha inválida genera incidencia.'),
('NUMERO','Texto','Número consignado en el registro.','Significado y unicidad pendientes; no se presume clave única.'),
('CANT','Número','Cantidad registrada de adquisición.','Interpretar junto con U#MEDIDA; no es variable objetivo.'),
('U#MEDIDA','Texto o categoría','Unidad de medida registrada.','Homologación pendiente; no sumar cantidades de unidades diferentes.'),
('DESCRIPCION','Texto','Descripción de lo adquirido.','Se homologa a descripcion y descripcion_normalizada; requiere catálogo aprobado.'),
('MONTO','Número monetario','Importe registrado de la adquisición.','Se homologa a importe_nominal en MXN; verificar ceros, moneda y alcance. No asumir CANT por P#UNITARIO sin conciliación.'),
('P#UNITARIO','Número monetario','Precio unitario registrado.','MXN por unidad consignada; definición y consistencia pendientes de revisión.'),
('CLASIFICACION','Categoría','Clasificación de origen.','No sustituye automáticamente al insumo_id aprobado.'),
('SUBCLASIFICACION','Categoría','Subclasificación de origen.','Homologar significado; uso auxiliar para revisar el catálogo.')]
source_info=[]
for fname,rows in [('Ventas.xlsx',sales),('Compras.xlsx',buys)]:
    f=Path('datasets/xlsx')/fname; w=load_workbook(f,read_only=True,data_only=True)
    headers=list(next(w['Hoja1'].values));w.close()
    assert headers==[r[0] for r in rows]
    source_info.append({'file':str(f),'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'fields':headers})

derived=[
('fecha','Fecha','FECHA de compras normalizada al día.','Clave temporal; solo registros previos al origen alimentan características.'),
('importe_nominal','Número en MXN','MONTO homologado en compras; importe semanal en contrato de ventas.','Objetivo monetario o predictor según tabla; inválidos no se sustituyen por cero.'),
('descripcion','Texto','DESCRIPCION homologada.','Base del catálogo, no una categoría aprobada por sí misma.'),
('descripcion_normalizada','Texto','Descripción normalizada mediante clean_upper.','Clave de enlace con catálogo; descripción vacía genera incidencia.'),
('registro_id','Entero','Número de fila de origen contado desde 2.','Trazabilidad dentro del archivo y su hash; no identificador universal.'),
('motivo','Texto','Resultado de revisión del registro.','Vacío significa sin incidencia detectada; no prueba cobertura completa.'),
('semana_inicio','Fecha','Lunes de la semana de lunes a domingo.','Clave semanal; agregación después de aprobar cobertura.'),
('insumo_id','Texto','Identificador asignado por catálogo aprobado.','No admite vacío ni nombres reservados total u otros.'),
('aprobado','Booleano','Aprobación de la correspondencia de catálogo.','El código admite true o 1; requiere decisión documentada.'),
('estado','Categoría','Condición de cobertura semanal.','observada o cero_confirmado para datos reales; desconocida o incompleta bloquean el intervalo. Ver nota de procedencia.'),
('evidencia','Texto','Sustento de la revisión de cobertura.','Obligatorio; debe justificar la clasificación semanal.'),
('fecha_revision','Fecha','Fecha de revisión de cobertura.','Obligatoria; no equivale a available_at de un predictor.'),
('registros_detectados','Entero','Conteo de filas por semana en plantilla.','Cero no confirma ausencia de compras; no establece cobertura.'),
('total','Número en MXN','Suma de importes semanales por insumo.','Objetivo monetario nominal; cero solo interpretable con cobertura validada.'),
('<insumo_id>','Número en MXN','Columnas dinámicas del panel, una por insumo.','Importe semanal agregado; nombres determinados por catálogo, no columnas literales fijas.'),
('otros','MXN o proporción','Residual de insumos no seleccionados, según matriz.','No confundir con Otros de ventas; distinguir matriz de importes de matriz de participaciones.'),
('Participación por insumo','Proporción de 0 a 1','Importe del insumo dividido entre total positivo.','Matriz con columnas por insumo; total cero implica no definido. Multiplicar por 100 solo para porcentaje.'),
('lag_1, lag_2, lag_4, lag_8, lag_12','Número en MXN','Totales de semanas previas al origen.','Predictores; sufijo indica distancia semanal. No usar valores objetivo futuros.'),
('media_4, media_8, media_12','Número en MXN','Media de las últimas semanas cerradas.','Ventana indicada por sufijo; requiere historia anterior al origen.'),
('desv_4, desv_8, desv_12','Número en MXN','Desviación estándar muestral de las semanas previas.','Mismas ventanas que las medias; un solo dato no define dispersión muestral.'),
('cal_sin, cal_cos','Número de −1 a 1','Codificación cíclica de semana objetivo.','Calendario conocido; periodo del código 52.1775 semanas.'),
('ventas_lag_1, ventas_lag_4','Número en MXN','Importe semanal de ventas rezagado.','Opcional; solo si available_at no es posterior al origen; de otro modo faltante.'),
('available_at','Fecha y hora','Momento de disponibilidad de ventas o exógenas.','Debe ser anterior o igual al origen. Ventas semanales: no antes del cierre.'),
('variable','Texto','Nombre de variable exógena.','Obligatorio en contrato opcional; no implica fuente incorporada.'),
('fecha_referencia','Fecha','Periodo al que corresponde una exógena.','No confundir con fecha de publicación o disponibilidad.'),
('valor','Número','Valor de variable exógena.','Finito; unidad específica a documentar por variable antes de incorporarla.'),
('fuente, version','Texto','Procedencia y versión de la exógena.','Obligatorias; permiten controlar revisiones de datos.'),
('tipo','Categoría','observada, pronostico o calendario.','Una observación futura no se etiqueta como disponible; pronósticos requieren archivo previo.'),
('exog_<variable>','Número','Valor exógeno admisible en el origen.','Pronóstico o calendario para la fecha objetivo; en su defecto última observación previa; ausente queda faltante.'),
('es_sintetico','Booleano','Indicador de procedencia sintética.','Control, nunca resultado real; la existencia del campo no autoriza usar esos datos.'),
('metodo_sintesis','Texto','Método declarado de generación sintética.','Metadato condicional del código; exige aprobación separada y no acredita observación real.')]

heading1=find('GLOSARIO'); heading2=next(p for p in ps if p.find('w:pPr/w:pStyle',ns) is not None and p.find('w:pPr/w:pStyle',ns).get(W+'val')=='Heading2') if any(p.find('w:pPr/w:pStyle',ns) is not None and p.find('w:pPr/w:pStyle',ns).get(W+'val')=='Heading2' for p in ps) else saleshead
annex=para('Anexo A Diccionario de datos',heading1)
ppr=annex.find('w:pPr',ns)
if ppr is None:ppr=E.SubElement(annex,W+'pPr')
E.SubElement(ppr,W+'pageBreakBefore')
annex_nodes=[annex,para('Este diccionario describe los campos de Ventas.xlsx y Compras.xlsx, hoja Hoja1, y las variables principales de entrada y preparación del pipeline inspeccionadas el 13 de septiembre de 2026. Los encabezados se verificaron directamente. Los tipos indicados son los esperados para el análisis, no una certificación de limpieza; las definiciones no confirmadas se señalan expresamente. No se reproducen datos personales ni valores de transacciones.'),
para('Reglas comunes de interpretación',heading2),
para('La existencia de un registro no prueba su integridad ni la cobertura de su semana. Cero, celda vacía y texto null tienen significados distintos: los valores cero se revisan con evidencia y los faltantes no se convierten automáticamente en ceros. Se observaron tipos mixtos y texto null en Ventas.xlsx. Antes de agregar se verificará si cada fila corresponde a una transacción, una línea o una repetición, y se documentará el tratamiento de devoluciones y cancelaciones.'),
para('Las fechas de inicio de captura provienen de la información histórica del negocio y no garantizan que los archivos inspeccionados contengan todo ese periodo. Los importes analíticos se expresarán en MXN nominales una vez confirmada la moneda y el alcance. Los campos administrativos y personales se excluyen del conjunto inicial de predictores. La disponibilidad temporal se comprobará antes de cada pronóstico; una fecha de operación no equivale por sí sola a fecha de disponibilidad.'),
para('Campos de origen de ventas',heading2),table(['Campo exacto','Tipo esperado','Definición','Uso y validación'],sales,[3.1,2.5,4.9,7]),
para('Campos de origen de compras',heading2),table(['Campo exacto','Tipo esperado','Definición','Uso y validación'],buys,[3.1,2.5,4.9,7]),
para('Variables analíticas y controles del pipeline',heading2),
para('Esta sección documenta contratos del código, no la aprobación de fuentes ni su incorporación efectiva. Los nombres entre signos angulares representan columnas dinámicas; los nombres agrupados identifican campos con la misma regla. Las participaciones se almacenan como proporciones y sus errores se reportan en puntos porcentuales. La inferencia futura no requiere importes observados de las semanas objetivo.'),
table(['Campo o familia','Tipo y unidad','Origen o regla','Uso y disponibilidad'],derived,[3.8,2.5,4.5,6.7]),
para('Nota de procedencia. El código inspeccionado contempla de manera condicional el estado sintetica_entrenamiento y campos de síntesis sujetos a aprobación y corte temporal. Su descripción no aprueba ni incorpora datos sintéticos al estudio. Esas observaciones no deben etiquetarse como reales ni utilizarse como evidencia de evaluación; cualquier diferencia entre el código y el protocolo científico deberá resolverse antes de la corrida definitiva.'),
para('Fuente. Elaboración propia a partir de los encabezados de datasets/xlsx/Ventas.xlsx y datasets/xlsx/Compras.xlsx, de los módulos data.py, features.py y composition.py del pipeline y de la información histórica proporcionada sobre Cup&Cake. Las equivalencias semánticas y unidades marcadas como pendientes requieren revisión de la documentación de origen antes de aprobar el conjunto analítico.')]
sect=body.find('w:sectPr',ns);position=list(body).index(sect) if sect is not None else len(body)
for i,n in enumerate(annex_nodes):body.insert(position+i,n)
assert before_eq==[E.tostring(n) for n in root.xpath('//m:oMath',namespaces=ns)]
assert method==[E.tostring(n) for n in list(body)[list(body).index(find('METODOLOGÍA')):list(body).index(find('DESARROLLO'))]]
assert refs==[E.tostring(n) for n in list(body)[list(body).index(find('REFERENCIAS')):list(body).index(annex)]]
backup=out/('Rev44_antes_diccionario_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.docx');backup.write_bytes(raw)
candidate=out/'Rev44_diccionario_verificada.docx'
with ZipFile(candidate,'w') as z:
    for info,b in entries:z.writestr(info,E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True) if info.filename=='word/document.xml' else b)
with ZipFile(candidate) as z:
    assert z.testzip() is None
    assert all(z.read(i.filename)==b for i,b in entries if i.filename!='word/document.xml')
assert path.read_bytes()==raw,'El documento cambió durante la edición'
Document(candidate)
path.write_bytes(candidate.read_bytes())
(out/'auditoria.json').write_text(json.dumps({'backup':str(backup),'sources':source_info,'source_fields':35,'analytical_rows':len(derived),'changes':changes,'methodology_preserved':True,'equations_preserved':len(before_eq),'references_preserved':True},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'document':str(path),'backup':str(backup),'source_fields':35,'analytical_rows':len(derived),'equations_preserved':len(before_eq)},ensure_ascii=False))
