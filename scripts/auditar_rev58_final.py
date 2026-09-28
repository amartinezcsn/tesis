from pathlib import Path
import re

from docx import Document


DOCX = Path(r"C:\Python\tesis\documentacion\TESIS_SEP2026_Rev58_(ZUJ)_Final.docx")
doc = Document(DOCX)

paragraphs = []
for index, paragraph in enumerate(doc.paragraphs, start=1):
    text = " ".join(paragraph.text.split())
    if text:
        paragraphs.append((index, paragraph.style.name, text))

print(f"paragraphs={len(paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")

targets = (
    "Pregunta de Investigación",
    "Objetivo general",
    "Objetivos específicos",
    "Hipótesis",
    "Tipo de investigación según el objetivo",
    "Diseño evaluativo",
    "Temporalidad de la investigación",
    "Validación temporal",
    "Consideraciones éticas",
    "Contraste de hipótesis",
    "Validez y limitaciones",
    "CONCLUSIONES",
    "Trabajo futuro",
)

for target in targets:
    for pos, (index, style, text) in enumerate(paragraphs):
        matches = text.casefold() == target.casefold() if target == "CONCLUSIONES" else target.casefold() in text.casefold()
        if index > 150 and matches:
            print(f"\n## {target} @párrafo {index} [{style}]")
            for row in paragraphs[pos : pos + 5]:
                print(f"{row[0]} [{row[1]}] {row[2]}")
            break

figures = []
tables = []
for index, style, text in paragraphs:
    match = re.match(r"^(Figura|Tabla)\s+(\d+)\b", text, re.I)
    if match and index > 150:
        (figures if match.group(1).casefold() == "figura" else tables).append(
            (int(match.group(2)), index, text)
        )

print("\nFIGURE_NUMBERS", [n for n, _, _ in figures])
print("TABLE_NUMBERS", [n for n, _, _ in tables])

in_refs = False
references = []
for index, style, text in paragraphs:
    if text == "REFERENCIAS":
        in_refs = True
        continue
    if in_refs and text.startswith("Anexo "):
        break
    if in_refs:
        references.append((index, text))

print(f"references={len(references)}")
for index, text in references:
    if "et al." in text or "DOI:" in text or "ResearchGate" in text:
        print(f"REFERENCE_FLAG {index}: {text}")

print("\nNUMBERED_HEADINGS")
for index, style, text in paragraphs:
    if re.match(r"^\d+(?:\.\d+){0,4}\.?\s+", text):
        print(f"{index} [{style}] {text}")
