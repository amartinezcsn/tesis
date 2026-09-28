from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Python\tesis\documentacion\TESIS_SEP2026_Rev58_(ZUJ)_Final.docx")
OUTPUT = Path(r"C:\Python\tesis\documentacion\TESIS_SEP2026_Rev59_(ZUJ)_Final.docx")


def paragraphs_everywhere(document: Document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def find_exact(document: Document, text: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip() == text.strip()]
    if len(matches) != 1:
        raise RuntimeError(f"Se esperaba una coincidencia para: {text[:100]!r}; encontradas={len(matches)}")
    return matches[0]


def set_exact(document: Document, old: str, new: str) -> None:
    find_exact(document, old).text = new


def replace_all(document: Document, old: str, new: str) -> int:
    count = 0
    for paragraph in paragraphs_everywhere(document):
        if old not in paragraph.text:
            continue
        # Most target text is held in one run. Preserve formatting whenever possible.
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                count += 1
                break
        else:
            paragraph.text = paragraph.text.replace(old, new)
            count += 1
    return count


def new_paragraph(text: str, style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph = Paragraph(element, None)
    if style:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def insert_paragraph_before(anchor: Paragraph, text: str, style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    if style:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def remove_numbering(paragraph: Paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is not None:
        ppr.remove(numpr)


def format_label(paragraph: Paragraph) -> None:
    paragraph.style = "Normal"
    remove_numbering(paragraph)
    paragraph.paragraph_format.space_before = Pt(9)
    paragraph.paragraph_format.space_after = Pt(3)
    for run in paragraph.runs:
        run.bold = True
        run.italic = True


def shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def style_table(table, widths: list[float] | None = None) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Inches(widths[c_idx])
            if r_idx == 0:
                shade_cell(cell, "D9EAF7")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 or c_idx != 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
                    if r_idx == 0:
                        run.bold = True


def append_table_after(document: Document, anchor: Paragraph, rows: list[list[str]], widths: list[float]):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for r_idx, values in enumerate(rows):
        for c_idx, value in enumerate(values):
            table.cell(r_idx, c_idx).text = value
    style_table(table, widths)
    anchor._p.addnext(table._tbl)
    return table


def add_update_fields_setting(document: Document) -> None:
    settings = document.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


doc = Document(SOURCE)
CHECK_DIR = Path(r"C:\Python\tesis\tmp\rev59_checkpoints")
CHECK_DIR.mkdir(parents=True, exist_ok=True)

# 1) Highest-impact correction: document and interpret the verified no-imputation sensitivity run.
set_exact(
    doc,
    "La evaluación produjo seis orígenes monetarios por horizonte y entre tres y cinco observaciones válidas por horizonte para la composición. El modelo híbrido obtuvo RMSE de 521.60, 695.41, 801.62 y 700.88 MXN para los horizontes de una a cuatro semanas, respectivamente, y solo alcanzó el menor RMSE entre los comparadores principales en el horizonte de tres semanas. En la distribución porcentual no existió un método dominante. Debido a la muestra reducida, 29 semanas de compras estimadas por promedio simple y la cobertura no certificada, el contraste de H1 fue no confirmatorio. El procedimiento se considera una aportación reproducible de carácter evaluativo exploratorio y un apoyo revisable para la planeación, no una orden automática de compra.",
    "La evaluación produjo seis orígenes monetarios por horizonte y entre tres y cinco observaciones válidas por horizonte para la composición. El modelo híbrido obtuvo RMSE de 521.60, 695.41, 801.62 y 700.88 MXN para los horizontes de una a cuatro semanas, respectivamente, y solo alcanzó el menor RMSE entre los comparadores principales en el horizonte de tres semanas. En la distribución porcentual no existió un método dominante. Un análisis de sensibilidad sin parámetros de imputación reprodujo exactamente métricas, selecciones, predicciones, particiones y contraste de H1; por tanto, las 29 semanas estimadas se conservaron solo como artefactos de auditoría y no intervinieron en el modelado. Debido a la muestra reducida y a la cobertura no certificada, el contraste de H1 fue no confirmatorio. El procedimiento constituye una aportación reproducible de carácter evaluativo exploratorio y un apoyo revisable para la planeación, no una orden automática de compra.",
)
set_exact(
    doc,
    "The evaluation produced six monetary forecast origins per horizon and between three and five valid composition observations per horizon. The hybrid model obtained RMSE values of MXN 521.60, 695.41, 801.62, and 700.88 for horizons one through four, respectively, and achieved the lowest RMSE among the main comparators only at the three-week horizon. No single percentage-allocation method dominated. Because of the small evaluation sample, 29 purchasing weeks estimated by simple averaging, and uncertified record coverage, the H1 assessment was non-confirmatory. The procedure is therefore presented as a reproducible exploratory evaluation and a revisable planning aid, not as an automated purchasing order.",
    "The evaluation produced six monetary forecast origins per horizon and between three and five valid composition observations per horizon. The hybrid model obtained RMSE values of MXN 521.60, 695.41, 801.62, and 700.88 for horizons one through four, respectively, and achieved the lowest RMSE among the main comparators only at the three-week horizon. No single percentage-allocation method dominated. A sensitivity run without imputation parameters reproduced the metrics, selections, predictions, partitions, and H1 assessment exactly; therefore, the 29 estimated weeks remained audit-only artifacts and did not enter model fitting or evaluation. Because of the small evaluation sample and uncertified record coverage, the H1 assessment was non-confirmatory. The procedure is presented as a reproducible exploratory evaluation and a revisable planning aid, not as an automated purchasing order.",
)

set_exact(
    doc,
    "Cuando se identifique una cola de semanas sin cobertura de compras, dichas semanas se trataron como ausencia de observación y no como importes cero. Cualquier extensión sintética se almacenó en un archivo independiente, con una etiqueta de procedencia, semilla y método de generación. Sus valores se reservaron únicamente en análisis de sensibilidad o pruebas técnicas del sistema de soporte a la decisión; se excluyeron del ajuste, la selección de modelos, la evaluación fuera de muestra y el contraste de H1.",
    "Las semanas sin cobertura de compras se trataron como ausencia de observación y no como importes cero. Las copias completadas por promedio simple se almacenaron en archivos derivados con su procedencia y método, pero el experimento consumió el panel original sin imputar. En consecuencia, esos valores se excluyeron del ajuste, la selección de modelos, la evaluación fuera de muestra y el contraste de H1.",
)
set_exact(
    doc,
    "La cobertura se determina con un calendario semanal y una tabla de evidencia revisada. Cada semana se distingue como observada, cero confirmado o desconocida; la ausencia de transacciones no se interpreta por sí sola como gasto cero. En la preparación acordada se completaron mediante promedio simple los huecos del tramo 2021-05-03 a 2024-04-29 para los archivos de ventas y compras de análisis. Esta serie completada se conserva como producto derivado y no sustituye los registros originales. Fuera del tramo autorizado, los periodos sin captura permanecen ausentes y se excluyen del ajuste cuando no hay evidencia suficiente.",
    "La cobertura se determinó con un calendario semanal y una tabla de evidencia revisada. Cada semana se distinguió como observada, cero confirmado o desconocida; la ausencia de transacciones no se interpretó por sí sola como gasto cero. En la preparación se generaron, como productos derivados de auditoría, copias completadas mediante promedio simple para el tramo del 3 de mayo de 2021 al 29 de abril de 2024. El experimento utilizó el panel original sin imputar; fuera del tramo autorizado, los periodos sin captura permanecieron ausentes y se excluyeron cuando no existió evidencia suficiente.",
)
set_exact(
    doc,
    "Los huecos del tramo autorizado del 3 de mayo de 2021 al 29 de abril de 2024 se completaron mediante promedio simple en copias de análisis de compras y ventas; cada valor estimado se identifica en los productos de auditoría. El resto de los periodos sin captura permanece ausente. Una semana sin transacciones no se convierte automáticamente en cero semanal: se conserva la diferencia entre cero respaldado y cobertura desconocida. Cinco líneas con importe $0 fechadas el 28 de agosto de 2023 se mantuvieron como registros y se exportaron por separado.",
    "Los huecos del tramo autorizado del 3 de mayo de 2021 al 29 de abril de 2024 se completaron mediante promedio simple únicamente en copias derivadas de auditoría; cada valor estimado quedó identificado. El panel original sin imputar alimentó el ajuste y la evaluación. El resto de los periodos sin captura permaneció ausente. Una semana sin transacciones no se convirtió automáticamente en cero semanal: se conservó la diferencia entre cero respaldado y cobertura desconocida. Cinco líneas con importe $0 fechadas el 28 de agosto de 2023 se mantuvieron como registros y se exportaron por separado.",
)
set_exact(
    doc,
    "El panel cubrió semanas entre el 5 de diciembre de 2016 y el 11 de mayo de 2026. Para el tramo de tres años destinado a completar series (3 de mayo de 2021–29 de abril de 2024), el proceso identificó 128 semanas de compras observadas y estimó 29 semanas faltantes mediante promedio simple; en ventas se reportaron 157 semanas observadas y ninguna semana estimada. Las estimaciones se conservaron identificadas como tales. Los huecos fuera de ese tramo no se rellenaron y se mantuvieron como ausencia de captura, por lo que no se interpretaron como gasto cero.",
    "El panel cubrió semanas entre el 5 de diciembre de 2016 y el 11 de mayo de 2026. En el tramo del 3 de mayo de 2021 al 29 de abril de 2024, la auditoría identificó 128 semanas de compras observadas y generó 29 semanas estimadas mediante promedio simple en un producto derivado; en ventas reportó 157 semanas observadas y ninguna estimada. Las 29 semanas no se incorporaron al panel utilizado para ajustar o evaluar modelos. Los huecos restantes se mantuvieron como ausencia de captura y no se interpretaron como gasto cero.",
)

design_heading = find_exact(doc, "Diseño efectivo de evaluación")
sens_heading = insert_paragraph_before(design_heading, "Análisis de sensibilidad de la imputación", "Heading 2")
sens_intro = insert_paragraph_before(
    design_heading,
    "Para verificar el posible efecto de las 29 semanas estimadas, se repitió el experimento con la misma semilla (42), fuentes, particiones y parámetros de modelado, pero sin los parámetros de imputación. La corrida principal run_20260924T195600Z_86a6f5 y la corrida de sensibilidad run_20260928T014913Z_38994f produjeron archivos idénticos, verificados por SHA-256, para métricas monetarias y composicionales, selección interna, predicciones, particiones y contraste de H1. La Tabla 5 resume la igualdad de las métricas monetarias por horizonte.",
    "Normal",
)
sens_caption = insert_paragraph_before(
    design_heading,
    "Tabla 5. Sensibilidad de las métricas monetarias a la configuración de imputación.",
    "Caption",
)
sens_rows = [
    ["Horizonte", "RMSE principal", "RMSE sin imputación", "MAE principal", "MAE sin imputación", "Diferencia"],
    ["h=1", "521.60", "521.60", "386.43", "386.43", "0.00"],
    ["h=2", "695.41", "695.41", "439.26", "439.26", "0.00"],
    ["h=3", "801.62", "801.62", "458.80", "458.80", "0.00"],
    ["h=4", "700.88", "700.88", "572.54", "572.54", "0.00"],
]
sens_table = doc.add_table(rows=len(sens_rows), cols=len(sens_rows[0]))
for r_idx, values in enumerate(sens_rows):
    for c_idx, value in enumerate(values):
        sens_table.cell(r_idx, c_idx).text = value
style_table(sens_table, [0.75, 1.05, 1.15, 1.0, 1.1, 0.8])
design_heading._p.addprevious(sens_table._tbl)
insert_paragraph_before(
    design_heading,
    "Nota. Errores en MXN; la diferencia corresponde a la corrida principal menos la corrida sin parámetros de imputación. La igualdad no demuestra que la imputación por promedio sea inocua en general: demuestra que, en esta implementación, las copias imputadas fueron productos de auditoría y no datos de entrenamiento o evaluación. Fuente: elaboración propia con las salidas de ambas corridas.",
    "Normal",
)

set_exact(
    doc,
    "La corrida demuestra que el flujo integra la limpieza y trazabilidad disponibles, produce comparaciones temporales a cuatro horizontes y genera una distribución que reconcilia con el total. En precisión, el desempeño favorable del híbrido dependió del horizonte y no dominó consistentemente a los métodos de referencia. La limitada cantidad de orígenes y de observaciones composicionales, las semanas sin captura, la imputación promedio en el tramo señalado y la falta de certificación de exhaustividad impiden generalizar el ranking o afirmar utilidad predictiva confirmatoria. La salida actual tampoco incorpora intervalos predictivos futuros.",
    "La corrida demuestra que el flujo integra la limpieza y trazabilidad disponibles, produce comparaciones temporales a cuatro horizontes y genera una distribución que reconcilia con el total. En precisión, el desempeño favorable del híbrido dependió del horizonte y no dominó consistentemente a los métodos de referencia. El análisis de sensibilidad confirmó que las copias imputadas no modificaron ninguna salida evaluativa. Aun así, la cantidad limitada de orígenes y observaciones composicionales, las semanas sin captura y la falta de certificación de exhaustividad impiden generalizar el ranking o afirmar utilidad predictiva confirmatoria. La salida actual tampoco incorpora intervalos predictivos futuros.",
)
set_exact(
    doc,
    "La imputación por promedio simple de 29 semanas de compras en el periodo especificado amplía la serie disponible, pero introduce valores estimados que suavizan la variabilidad y no reemplazan comprobantes faltantes. Aunque se mantuvieron separados de las observaciones originales, esta decisión forma parte del contexto de los resultados. La falta de certificación de captura exhaustiva impide asegurar que el objetivo represente la totalidad de las compras de cada semana; las conclusiones corresponden a importes registrados utilizables y a las categorías incluidas.",
    "La auditoría generó 29 semanas estimadas por promedio simple en un producto derivado, pero la inspección del flujo y la corrida de sensibilidad confirmaron que esos valores no entraron al ajuste, a la selección ni a la evaluación: las salidas clave fueron idénticas con y sin la configuración de imputación. Por ello, no se atribuye a esas estimaciones un sesgo en las métricas reportadas. La limitación que permanece es la cobertura original: la falta de certificación de captura exhaustiva impide asegurar que el objetivo represente la totalidad de las compras de cada semana; las conclusiones corresponden a importes registrados utilizables y a las categorías incluidas.",
)
set_exact(
    doc,
    "Respecto de los objetivos específicos, se depuraron y agregaron semanalmente las fuentes disponibles y se caracterizó el panel, manteniendo identificadas 29 semanas de compras estimadas por promedio simple en el tramo establecido y conservando los huecos restantes como ausencia de captura. Se implementaron combinaciones estadísticas y de aprendizaje automático, con información histórica de ventas y calendario disponible al origen; no se aisló mediante una prueba de ablación cuánto aporta cada predictor. También se estimaron y evaluaron participaciones históricas y ponderadas en el tiempo: su desempeño alternó entre horizontes, sin un método dominante en todos ellos. Finalmente, el contraste temporal frente a referencias se completó en carácter exploratorio, condicionado por la muestra reducida y la cobertura no certificada.",
    "Respecto de los objetivos específicos, se depuraron y agregaron semanalmente las fuentes disponibles y se caracterizó el panel. Las 29 semanas estimadas por promedio simple se mantuvieron identificadas en productos derivados de auditoría y se excluyeron del modelado; la repetición sin parámetros de imputación produjo exactamente las mismas salidas clave. Se implementaron combinaciones estadísticas y de aprendizaje automático, con información histórica de ventas y calendario disponible al origen; no se aisló mediante una prueba de ablación cuánto aporta cada predictor. También se estimaron y evaluaron participaciones históricas y ponderadas en el tiempo: su desempeño alternó entre horizontes, sin un método dominante. El contraste temporal frente a referencias se completó con carácter exploratorio, condicionado por la muestra reducida y la cobertura no certificada.",
)
set_exact(
    doc,
    "Antes de emplear el procedimiento para planear compras, se recomienda completar y documentar la auditoría de cobertura, resolver discrepancias de captura y revisar las semanas completadas por promedio simple con los registros primarios disponibles. Deben mantenerse trazables las categorías excluidas y la categoría RESTO_ELEGIBLE, que concilia insumos elegibles menores y no representa OTROS. Cada emisión debe indicar fecha de corte, periodo objetivo, horizonte, supuestos y carácter nominal de los importes.",
    "Antes de emplear el procedimiento para planear compras, se recomienda completar y documentar la auditoría de cobertura y resolver discrepancias de captura con los registros primarios disponibles. Las copias imputadas deben conservarse separadas y trazables como productos de auditoría, sin incorporarse al entrenamiento o a la evaluación salvo que un protocolo futuro lo autorice y lo analice expresamente. También deben mantenerse trazables las categorías excluidas y RESTO_ELEGIBLE, que concilia insumos elegibles menores y no representa OTROS. Cada emisión debe indicar fecha de corte, periodo objetivo, horizonte, supuestos y carácter nominal de los importes.",
)
set_exact(
    doc,
    "El trabajo posterior deberá priorizar la ampliación de registros primarios de compras y ventas para reducir la proporción imputada y aumentar los periodos evaluables. Con una base más completa, será posible realizar una evaluación confirmatoria con suficientes orígenes y estimar intervalos predictivos para el importe total y las asignaciones por insumo. También será pertinente evaluar por separado el aporte de ventas, calendario y otras variables exógenas mediante experimentos de ablación temporal.",
    "El trabajo posterior deberá priorizar la ampliación y certificación de los registros primarios de compras y ventas para reducir los huecos y aumentar los periodos evaluables. Con una base más completa, será posible realizar una evaluación confirmatoria con suficientes orígenes y estimar intervalos predictivos para el importe total y las asignaciones por insumo. También será pertinente evaluar por separado el aporte de ventas, calendario y otras variables exógenas mediante experimentos de ablación temporal.",
)
doc.save(CHECK_DIR / "01_sensibilidad.docx")

# 2) Correct verb tense for the completed study while preserving future recommendations.
tense_replacements = {
    "Como regla previa del estudio, H1 requerirá evidencia favorable": "Como regla previa del estudio, H1 requería evidencia favorable",
    "la investigación deberá documentar": "la investigación documentó",
    "El pronóstico tendrá como propósito": "El pronóstico tuvo como propósito",
    "evaluó herramientas de pronóstico semanal con la información disponible en una microempresa y valorará su pertinencia": "evaluó herramientas de pronóstico semanal con la información disponible en una microempresa y valoró su pertinencia",
    "el estudio buscará determinar": "el estudio buscó determinar",
    "Los resultados podrían aportar": "Los resultados aportaron",
    "el procedimiento podrá servir": "el procedimiento puede servir",
    "su aplicación requerirá evaluar": "su aplicación requiere evaluar",
    "Se propone aprovechar": "Se aprovecharon",
    "La contribución prevista será": "La contribución fue",
    "y podrá orientar": "y puede orientar",
    "que permita consultar": "que permite consultar",
    "No generará órdenes": "No genera órdenes",
    "estas funciones requerirían": "estas funciones requieren",
    "todavía deberá contrastarse": "se contrastó",
    "no se adoptará el alcance causal": "no se adoptó el alcance causal",
    "La aplicación tendrá carácter": "La aplicación tuvo carácter",
    "se contrastarán alternativas": "se contrastaron alternativas",
    "y comparará la precisión": "y comparó la precisión",
    "Se calculón el importe": "Se calculó el importe",
    "cuando su información esté disponible": "cuando su información estaba disponible",
    "no implicará generalización": "no implicó generalización",
    "No se introducirán intervenciones": "No se introdujeron intervenciones",
    "ni se asignarán grupos": "ni se asignaron grupos",
    "La evaluación reproducirá situaciones": "La evaluación reprodujo situaciones",
    "y utilizará exclusivamente": "y utilizó exclusivamente",
    "la referencia primaria será": "la referencia primaria fue",
    "cuando exista cobertura suficiente": "cuando existió cobertura suficiente",
    "se concretará mediante": "se concretó mediante",
    "Los registros se organizarán": "Los registros se organizaron",
    "se documentarán antes": "se documentaron antes",
    "se contrastará con su propia": "se contrastó con su propia",
    "tendrán carácter complementario": "tuvieron carácter complementario",
    "no constituirán una segunda": "no constituyeron una segunda",
    "La evidencia favorable deberá abarcar": "La evidencia favorable debía abarcar",
    "comunicará este inventario": "comunicó este inventario",
    "se resolverán desde": "se resolvieron desde",
    "Este diccionario documentará": "Este diccionario documentó",
    "no se utilizón": "no se utilizaron",
    "Su función será fundamentar": "Su función fue fundamentar",
    "se ajustará exclusivamente": "se ajustó exclusivamente",
    "utilizará un criterio propio": "utilizó un criterio propio",
    "no se combinarán directamente": "no se combinaron directamente",
    "no tendrán una composición": "no tuvieron una composición",
    "cuando correspondan a ceros": "cuando correspondieron a ceros",
}
for old, new in tense_replacements.items():
    replace_all(doc, old, new)
doc.save(CHECK_DIR / "02_tiempos.docx")

# 3) Correct bibliographic attribution and APA-visible issues.
replace_all(doc, "Ab Karim et al. (2020)", "Lee et al. (2020)")
replace_all(doc, "(Ab Karim et al., 2020)", "(Lee et al., 2020)")
replace_all(doc, "Aprodu et al. (2025)", "Melesse y Orrù (2025)")
replace_all(doc, "(Aprodu et al., 2025)", "(Melesse & Orrù, 2025)")
set_exact(
    doc,
    "Ab Karim, M. S., Ab Rahman, S., Chua, B. L., et al. (2020). The creative minds of extraordinary pastry chefs: an integrated theory of aesthetic expressions – a portraiture study. International Journal of Contemporary Hospitality Management, 32(9), 3015–3034. https://doi.org/10.1108/IJCHM-04-2020-0329",
    "Lee, K.-S., Blum, D., Miao, L., & Tomas, S. R. (2020). The creative minds of extraordinary pastry chefs: An integrated theory of aesthetic expressions—A portraiture study. International Journal of Contemporary Hospitality Management, 32(9), 3015–3034. https://doi.org/10.1108/IJCHM-04-2020-0329",
)
set_exact(
    doc,
    "Aprodu, I., Banu, I., Vasilean, I., et al. (2025). The Digital Revolution in the Bakery Sector: A Systematic Review of Industry 4.0 Technologies. Foods, 14(2).",
    "Melesse, T. Y., & Orrù, P. F. (2025). The digital revolution in the bakery sector: Innovations, challenges, and opportunities from Industry 4.0. Foods, 14(3), Article 526. https://doi.org/10.3390/foods14030526",
)
set_exact(
    doc,
    "Ben Abdallah, H., Ayadi, M., Boujelbene, Y. (2022). The main determinants and effects of product innovation: An exploratory study on the pastry companies of the region of Sfax (in Tunisia). Technological Forecasting and Social Change, 185, 122065. DOI: 10.1016/j.techfore.2022.122065.",
    "Ben Abdallah, H., Ayadi, M., & Boujelbene, Y. (2022). The main determinants and effects of product innovation: An exploratory study on the pastry companies of the region of Sfax (in Tunisia). Technological Forecasting and Social Change, 185, Article 122065. https://doi.org/10.1016/j.techfore.2022.122065",
)

# Keep references alphabetized after the corrected first author names.
refs_heading = find_exact(doc, "REFERENCIAS")
appendix_heading = find_exact(doc, "Anexo A Diccionario de datos")
ref_elements = []
node = refs_heading._p.getnext()
while node is not None and node is not appendix_heading._p:
    following = node.getnext()
    if node.tag == qn("w:p"):
        ref_elements.append(node)
    node = following
for element in ref_elements:
    element.getparent().remove(element)
for element in sorted(ref_elements, key=lambda e: "".join(e.itertext()).casefold()):
    appendix_heading._p.addprevious(element)
doc.save(CHECK_DIR / "03_referencias.docx")

# 4) Sequential table/figure numbering and cross-references.
figure_map = {33: 1, 34: 2, 35: 3, 36: 4, 37: 5, 38: 6, 39: 7, 40: 8, 41: 9, 42: 10, 43: 11, 46: 12, 47: 13, 48: 14, 49: 15}
for old, new in figure_map.items():
    replace_all(doc, f"Figura {old}", f"Figura {new}")
for old, new in {12: 2, 13: 3, 14: 4}.items():
    replace_all(doc, f"Tabla {old}", f"Tabla {new}")
replace_all(doc, "Tabla 1 -", "Tabla 1.")

# Remove unintended multilevel numbering from descriptive sublabels.
for label in (
    "Objetivo general",
    "Objetivos específicos",
    "Relevancia social",
    "Conveniencia",
    "Valor teórico",
    "Implicaciones prácticas",
    "Hipótesis de investigación (H1)",
    "Reglas comunes de interpretación",
    "Campos de origen de ventas",
    "Campos de origen de compras",
    "Variables analíticas y controles del pipeline",
):
    matches = [p for p in doc.paragraphs if p.text.strip() == label]
    if len(matches) != 1:
        raise RuntimeError(f"Etiqueta no única: {label!r}, encontradas={len(matches)}")
    format_label(matches[0])
doc.save(CHECK_DIR / "04_numeracion.docx")

# Complete-study wording in Appendix A.
replace_all(doc, "Antes de agregar se verificará", "Antes de agregar se verificó")
replace_all(doc, "se documentará el tratamiento", "se documentó el tratamiento")
replace_all(doc, "Los importes analíticos se expresarán", "Los importes analíticos se expresaron")
replace_all(doc, "La disponibilidad temporal se comprobará", "La disponibilidad temporal se comprobó")
replace_all(doc, "deberá resolverse antes de la corrida definitiva", "se resolvió antes de la corrida evaluada")

# 5) Add an explicit reproducibility appendix.
appendix_b = doc.add_paragraph("Anexo B Paquete de reproducibilidad", style="Heading 1")
appendix_b.paragraph_format.page_break_before = True
doc.add_paragraph(
    "El paquete se conserva dentro de la raíz del proyecto de tesis. Las rutas siguientes son relativas a esa raíz y permiten localizar la configuración, el código y los productos que sustentan los resultados. Los datos transaccionales originales permanecen restringidos por confidencialidad; los manifiestos registran huellas SHA-256, versiones de bibliotecas, parámetros y semilla aleatoria (42)."
)
package_rows = [
    ["Componente", "Ruta relativa", "Función"],
    ["Entrada del pipeline", "codigos/00_pipeline_hibrido.py", "Punto de ejecución reproducible"],
    ["Dependencias", "codigos/requirements_hibrido.txt", "Versiones requeridas"],
    ["Configuración principal", "codigos/02_config_hibrido.json", "Parámetros de la corrida evaluada"],
    ["Configuración de sensibilidad", "codigos/02_config_hibrido_sensibilidad_sin_imputacion_rev59.json", "Mismos parámetros, sin imputación"],
    ["Manifiesto principal", "output/run_20260924T195600Z_86a6f5/manifiesto.json", "Huellas, entorno y parámetros"],
    ["Manifiesto de sensibilidad", "output/run_20260928T014913Z_38994f/manifiesto.json", "Trazabilidad de la repetición"],
    ["Métricas y selección", "output/run_20260924T195600Z_86a6f5/metricas.csv; seleccion_interna.csv", "Resultados monetarios y elección interna"],
    ["Composición", "output/run_20260924T195600Z_86a6f5/metricas_composicion.csv; seleccion_composicion.csv", "Resultados de participaciones"],
    ["Predicciones y particiones", "output/run_20260924T195600Z_86a6f5/predicciones.csv; particiones.csv", "Observaciones, pronósticos y cortes"],
    ["Contraste de H1", "output/run_20260924T195600Z_86a6f5/hipotesis.json", "Resultado del contraste exploratorio"],
]
package_table = doc.add_table(rows=len(package_rows), cols=3)
for r_idx, values in enumerate(package_rows):
    for c_idx, value in enumerate(values):
        package_table.cell(r_idx, c_idx).text = value
style_table(package_table, [1.35, 3.55, 1.75])
doc.add_paragraph(
    "Criterio de reproducción. Una repetición debe conservar las fuentes autorizadas, la configuración, la semilla y las versiones registradas en el manifiesto. Para la sensibilidad se suprimieron únicamente los parámetros de imputación. La identidad SHA-256 de metricas.csv, metricas_composicion.csv, seleccion_interna.csv, seleccion_composicion.csv, predicciones.csv, particiones.csv e hipotesis.json entre ambas corridas demuestra que las copias imputadas no intervinieron en las salidas evaluativas."
)
doc.add_paragraph(
    "Restricción de acceso. Compras.xlsx, Ventas.xlsx y los archivos contextuales contienen información operativa de la empresa y no se incorporan al documento. Su custodia corresponde al proyecto autorizado; la reproducción por terceros requiere acceso legítimo a las mismas fuentes o a una versión anonimizada aprobada."
)
doc.save(CHECK_DIR / "05_anexo.docx")

# Continuous Arabic pagination: keep the first numbered front-matter restart, remove the Abstract restart.
for idx, section in enumerate(doc.sections):
    pg_num = section._sectPr.find(qn("w:pgNumType"))
    if idx >= 2 and pg_num is not None and qn("w:start") in pg_num.attrib:
        del pg_num.attrib[qn("w:start")]
        if not pg_num.attrib and len(pg_num) == 0:
            section._sectPr.remove(pg_num)

# Add placeholders to the static contents list; page numbers are finalized after pagination.
toc_appendix_a = next(p for p in doc.paragraphs if p.style.name.lower() == "toc 1" and p.text.startswith("Anexo A"))
toc_b = OxmlElement("w:p")
toc_appendix_a._p.addnext(toc_b)
toc_b_p = Paragraph(toc_b, toc_appendix_a._parent)
toc_b_p.style = "toc 1"
toc_b_p.add_run("Anexo B Paquete de reproducibilidad\tPENDIENTE")
toc_results = next(p for p in doc.paragraphs if p.style.name.lower() == "toc 2" and p.text.startswith("Calidad, cobertura"))
toc_sens = OxmlElement("w:p")
toc_results._p.addnext(toc_sens)
toc_sens_p = Paragraph(toc_sens, toc_results._parent)
toc_sens_p.style = "toc 2"
toc_sens_p.add_run("Análisis de sensibilidad de la imputación\tPENDIENTE")

# Replace the static index with provisional complete entries. Final page numbers are populated after render.
figure_index_heading = find_exact(doc, "ÍNDICE DE TABLAS Y FIGURAS")
summary_heading = find_exact(doc, "RESUMEN")
node = figure_index_heading._p.getnext()
while node is not None and node is not summary_heading._p:
    following = node.getnext()
    node.getparent().remove(node)
    node = following

caption_texts = []
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if re.match(r"^(Tabla|Figura)\s+\d+[.\-]", text, re.I) and not paragraph.style.name.lower().startswith("table of figures"):
        caption_texts.append(text)
anchor = figure_index_heading
for text in caption_texts:
    p_el = OxmlElement("w:p")
    anchor._p.addnext(p_el)
    p = Paragraph(p_el, anchor._parent)
    p.style = "table of figures"
    p.add_run(f"{text}\tPENDIENTE")
    anchor = p
doc.save(CHECK_DIR / "06_indices.docx")

doc.core_properties.comments = (
    "Rev59: sensibilidad de imputación verificada, referencias corregidas, numeración secuencial, "
    "tiempos verbales ajustados, paginación arábiga continua y anexo de reproducibilidad."
)
doc.save(OUTPUT)

print(f"saved={OUTPUT}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
print(f"captions={len(caption_texts)}")
