"""Aplica el ajuste 2A a la Rev41 sin modificar el documento fuente."""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(r"C:\Python\tesis")
SOURCE = ROOT / "documentacion" / "TESIS_AGO2026_Rev41_(ZUJ)_29ago2026_AnalisisSeriesTemporales.docx"
OUTPUT = ROOT / "documentacion" / "TESIS_AGO2026_Rev42_(ZUJ)_05sep2026_DelimitacionCientifica.docx"

TITLE = (
    "PRONÓSTICO DEL IMPORTE SEMANAL DE COMPRAS PARA EL PRESUPUESTO DE "
    "ABASTECIMIENTO MEDIANTE MODELOS ESTADÍSTICOS Y APRENDIZAJE AUTOMÁTICO "
    "EN UNA MICROEMPRESA DE REPOSTERÍA DE TIZAYUCA HIDALGO"
)


def find_exact(document: Document, text: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == text.strip():
            return paragraph
    raise ValueError(f"No se encontró el párrafo: {text}")


def find_startswith(document: Document, text: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(text):
            return paragraph
    raise ValueError(f"No se encontró un párrafo que inicie con: {text}")


def replace_text(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def paragraph_after(document: Document, anchor, text: str = "", style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    if text:
        paragraph.add_run(text)
    anchor._p.addnext(paragraph._p)
    return paragraph


def math_run(text: str, *, upright: bool = False):
    run = OxmlElement("m:r")
    if upright:
        properties = OxmlElement("m:rPr")
        style = OxmlElement("m:sty")
        style.set(qn("m:val"), "p")
        properties.append(style)
        run.append(properties)
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    return run


def subscript(base_nodes, sub_nodes):
    element = OxmlElement("m:sSub")
    element.append(OxmlElement("m:sSubPr"))
    base = OxmlElement("m:e")
    for node in base_nodes:
        base.append(node)
    sub = OxmlElement("m:sub")
    for node in sub_nodes:
        sub.append(node)
    element.extend((base, sub))
    return element


def subscript_superscript(base_nodes, sub_nodes, super_nodes):
    element = OxmlElement("m:sSubSup")
    element.append(OxmlElement("m:sSubSupPr"))
    base = OxmlElement("m:e")
    for node in base_nodes:
        base.append(node)
    sub = OxmlElement("m:sub")
    for node in sub_nodes:
        sub.append(node)
    sup = OxmlElement("m:sup")
    for node in super_nodes:
        sup.append(node)
    element.extend((base, sub, sup))
    return element


def target_equation():
    equation = OxmlElement("m:oMath")
    equation.append(subscript([math_run("y")], [math_run("t")]))
    equation.append(math_run(" = ", upright=True))

    summation = OxmlElement("m:nary")
    properties = OxmlElement("m:naryPr")
    character = OxmlElement("m:chr")
    character.set(qn("m:val"), "∑")
    limit_location = OxmlElement("m:limLoc")
    limit_location.set(qn("m:val"), "undOvr")
    hide_super = OxmlElement("m:supHide")
    hide_super.set(qn("m:val"), "1")
    properties.extend((character, limit_location, hide_super))
    summation.append(properties)

    lower = OxmlElement("m:sub")
    lower.append(math_run("d"))
    lower.append(math_run(" ∈ ", upright=True))
    lower.append(subscript([math_run("W")], [math_run("t")]))
    summation.append(lower)
    summation.append(OxmlElement("m:sup"))

    expression = OxmlElement("m:e")
    expression.append(
        subscript_superscript(
            [math_run("C")],
            [math_run("d")],
            [math_run("real", upright=True)],
        )
    )
    summation.append(expression)
    equation.append(summation)
    return equation


def forecast_equation():
    accent = OxmlElement("m:acc")
    accent_properties = OxmlElement("m:accPr")
    character = OxmlElement("m:chr")
    character.set(qn("m:val"), "̂")
    accent_properties.append(character)
    accent.append(accent_properties)
    base = OxmlElement("m:e")
    base.append(math_run("y"))
    accent.append(base)

    equation = OxmlElement("m:oMath")
    equation.append(
        subscript(
            [accent],
            [math_run("t+h|t")],
        )
    )
    equation.append(math_run(",    h ∈ {1, 2, 3, 4}", upright=True))
    return equation


def equation_after(document: Document, anchor, equation):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    math_paragraph = OxmlElement("m:oMathPara")
    math_properties = OxmlElement("m:oMathParaPr")
    justification = OxmlElement("m:jc")
    justification.set(qn("m:val"), "centerGroup")
    math_properties.append(justification)
    math_paragraph.extend((math_properties, equation))
    paragraph._p.append(math_paragraph)
    anchor._p.addnext(paragraph._p)
    return paragraph


def main() -> None:
    document = Document(SOURCE)

    title = document.paragraphs[3]
    replace_text(title, TITLE)
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(0)
    title_properties = title._p.get_or_add_pPr()
    border = title_properties.find(qn("w:pBdr"))
    if border is not None:
        title_properties.remove(border)
    for run in title.runs:
        run.bold = True
        run.font.name = "Arial"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)

    first_scope = find_startswith(document, "El reto de Cup&Cake es fortalecer")
    second_scope = find_startswith(document, "El pronóstico tendrá como propósito")

    heading = first_scope.insert_paragraph_before(
        "Delimitación científica y unidad de análisis",
        style="Heading 4",
    )
    heading.paragraph_format.keep_with_next = True

    replace_text(
        first_scope,
        "El objeto predictivo de esta investigación es el importe real de compras por semana calendario. "
        "Si Wₜ representa el conjunto de días de la semana t y C el importe diario de compras "
        "expresado en pesos reales de mayo de 2026, la variable objetivo se define como:",
    )
    first_scope.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    target_math = equation_after(document, first_scope, target_equation())
    definition = paragraph_after(
        document,
        target_math,
        "En esta expresión, yₜ es el importe semanal observado. La definición fija una unidad de análisis "
        "monetaria y semanal. Los registros de ventas, calendario, INPC y temperatura se consideran "
        "predictores sólo cuando estaban disponibles antes del periodo pronosticado.",
    )
    definition.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    replace_text(
        second_scope,
        "Desde cada origen temporal t se estiman de forma directa los importes de las semanas futuras. "
        "La notación del pronóstico y los horizontes evaluados se expresan como:",
    )
    second_scope.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    forecast_math = equation_after(document, second_scope, forecast_equation())
    horizon_text = paragraph_after(
        document,
        forecast_math,
        "El horizonte h=1 constituye la evaluación principal y h=4 aporta evidencia complementaria. "
        "El presupuesto de cuatro semanas se obtiene sumando los pronósticos directos de h=1, h=2, h=3 "
        "y h=4; por tanto, no se interpreta como un modelo mensual independiente.",
    )
    horizon_text.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    limitation = paragraph_after(
        document,
        horizon_text,
        "El propósito inferencial se limita a cuantificar la precisión predictiva fuera de muestra y a "
        "apoyar la revisión del presupuesto de abastecimiento. El diseño no permite atribuir efectos "
        "causales sobre inventarios, merma, rentabilidad, expansión comercial o desabasto, ni genera "
        "órdenes automáticas o cantidades físicas de compra.",
    )
    limitation.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    document.core_properties.title = TITLE.title()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
