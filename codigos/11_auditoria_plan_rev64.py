from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path(r"C:\Python\tesis\documentacion")
FILES = {
    "rev62": ROOT / "TESIS_SEP2026_Rev62_(ZUJ)_Final.docx",
    "rev63": ROOT / "TESIS_OCT2026_Rev63_(ZUJ)_Revision_Comite.docx",
    "rev64": ROOT / "TESIS_OCT2026_Rev64_(ZUJ)_Correcciones_Director.docx",
}
OUT = Path(r"C:\Python\tesis\tmp\inventario_rev64\auditoria_estructural.json")


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def has_numpr(paragraph) -> bool:
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.numPr is not None


def next_nonempty(paragraphs, start, limit=6):
    result = []
    for p in paragraphs[start + 1 : start + 1 + limit]:
        if norm(p.text):
            result.append(norm(p.text))
    return result


def audit(path: Path):
    doc = Document(path)
    paragraphs = list(doc.paragraphs)
    headings = []
    for i, p in enumerate(paragraphs):
        style = p.style.name if p.style else ""
        if style.lower().startswith(("heading", "título", "titulo")):
            headings.append({
                "index": i,
                "style": style,
                "text": norm(p.text),
                "numbering_property": has_numpr(p),
            })

    captions = []
    for i, p in enumerate(paragraphs):
        style = p.style.name if p.style else ""
        text = norm(p.text)
        if style.lower() == "caption" or re.match(r"^(Figura|Tabla)\b", text, re.I):
            captions.append({
                "index": i,
                "text": text,
                "alignment": str(p.alignment),
                "centered": p.alignment == WD_ALIGN_PARAGRAPH.CENTER,
            })

    acknowledgements = []
    for i, p in enumerate(paragraphs):
        if norm(p.text).upper() == "AGRADECIMIENTOS":
            acknowledgements = next_nonempty(paragraphs, i, 8)

    reference_lines = []
    in_refs = False
    for p in paragraphs:
        text = norm(p.text)
        if text.upper() == "REFERENCIAS":
            in_refs = True
            continue
        if in_refs and (p.style and p.style.name.lower().startswith("heading 1")):
            break
        if in_refs and text:
            reference_lines.append(text)

    figure_captions = [c for c in captions if c["text"].lower().startswith("figura")]
    table_captions = [c for c in captions if c["text"].lower().startswith("tabla")]
    return {
        "file": str(path),
        "paragraphs": len(paragraphs),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "heading1": [h for h in headings if h["style"].lower().endswith("1")],
        "numbered_headings": sum(1 for h in headings if h["numbering_property"]),
        "headings_total": len(headings),
        "acknowledgements_following": acknowledgements,
        "reference_paragraph_count": len(reference_lines),
        "reference_preview": reference_lines[:3],
        "figure_caption_count": len(figure_captions),
        "figure_captions_centered": sum(c["centered"] for c in figure_captions),
        "table_caption_count": len(table_captions),
        "table_captions_centered": sum(c["centered"] for c in table_captions),
        "caption_examples": captions[:5],
    }


result = {name: audit(path) for name, path in FILES.items()}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
safe = json.dumps(result, ensure_ascii=False, indent=2).encode("cp1252", "replace").decode("cp1252")
print(safe)
