from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Python\tesis")
SRC = ROOT / "documentacion" / "TESIS_SEP2026_Rev62_(ZUJ)_Final.docx"
OUT = ROOT / "documentacion" / "TESIS_OCT2026_Rev64_(ZUJ)_Correcciones_Director.docx"


def norm(text: str) -> str:
    return " ".join(text.split())


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def replace_text(paragraph, text: str):
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def remove_numbering(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is not None:
        ppr.remove(numpr)


def find_paragraph(doc, text: str, exact=True):
    for paragraph in doc.paragraphs:
        current = norm(paragraph.text)
        if (exact and current == text) or (not exact and text in current):
            return paragraph
    raise ValueError(f"No se encontró el párrafo: {text}")


def paragraph_index(doc, paragraph):
    for index, candidate in enumerate(doc.paragraphs):
        if candidate._p is paragraph._p:
            return index
    raise ValueError("El párrafo ya no pertenece al cuerpo del documento")


def insert_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p.getparent().remove(new_para._p)
    new_p.getparent().replace(new_p, new_para._p)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def insert_before(paragraph, text="", style=None):
    new_para = insert_after(paragraph, text, style)
    paragraph._p.addprevious(new_para._p)
    return new_para


def move_range_after(start, end_exclusive, target):
    nodes = []
    node = start._p
    while node is not end_exclusive._p:
        nxt = node.getnext()
        nodes.append(node)
        node = nxt
    anchor = target._p
    for item in nodes:
        anchor.addnext(item)
        anchor = item


def set_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    header = trpr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        trpr.append(header)
    header.set(qn("w:val"), "true")


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, widths):
    table.style = "Table Grid"
    table.autofit = False
    for i, row in enumerate(table.rows):
        if i == 0:
            set_repeat_header(row)
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width = Inches(widths[j])
            set_cell_margins(cell)
            if i == 0:
                shade_cell(cell, "1F4E79")
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(8.5)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                if i % 2 == 0:
                    shade_cell(cell, "EAF2F8")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.size = Pt(8)


def add_bookmarked_caption(paragraph, label, number, title, bookmark_id):
    clear_paragraph(paragraph)
    paragraph.style = "Caption"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), f"_Ref{label}{number}")
    paragraph._p.append(start)
    paragraph.add_run(f"{label} {number}. {title}")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.append(end)


def insert_table_block(after_paragraph, number, title, headers, rows, widths, intro, note, bookmark_id):
    intro_p = insert_after(after_paragraph, intro, "Normal")
    intro_p.paragraph_format.space_before = Pt(6)
    caption = insert_after(intro_p, style="Caption")
    add_bookmarked_caption(caption, "Tabla", number, title, bookmark_id)
    table = doc.add_table(rows=1, cols=len(headers))
    for j, header in enumerate(headers):
        table.rows[0].cells[j].text = header
    for values in rows:
        cells = table.add_row().cells
        for j, value in enumerate(values):
            cells[j].text = value
    style_table(table, widths)
    caption._p.addnext(table._tbl)
    note_p = insert_after(caption, note, "Normal")
    table._tbl.addnext(note_p._p)
    note_p.paragraph_format.space_after = Pt(6)
    return note_p


doc = Document(SRC)

# 1. Preliminares y agradecimientos.
ack = find_paragraph(doc, "AGRADECIMIENTOS")
ack_next = doc.paragraphs[paragraph_index(doc, ack) + 1]
replace_text(
    ack_next,
    "A la Universidad Politécnica Metropolitana de Hidalgo, por brindar el espacio académico para desarrollar esta investigación. "
    "A la dirección y al comité de tesis, por sus observaciones metodológicas y editoriales. A Cup&Cake, por autorizar el uso académico "
    "de sus registros y facilitar la comprensión del contexto operativo. Finalmente, a mi familia y a las personas que acompañaron este proceso, "
    "por su apoyo durante la elaboración del trabajo.",
)
for paragraph in list(doc.paragraphs):
    if norm(paragraph.text) == "A la Universidad Politécnica Metropolitana de Hidalgo":
        paragraph._p.getparent().remove(paragraph._p)

index_fig = find_paragraph(doc, "ÍNDICE DE IMÁGENES")
replace_text(index_fig, "ÍNDICE DE FIGURAS")

# 2. Reorganización en cinco capítulos. El estado del arte se traslada al inicio del Capítulo II.
intro = find_paragraph(doc, "INTRODUCCIÓN")
replace_text(intro, "CAPÍTULO I INTRODUCCIÓN")
remove_numbering(intro)

antecedentes = insert_after(intro, "1.1 Antecedentes", "Heading 2")
remove_numbering(antecedentes)

state_start = find_paragraph(doc, "Estado del Arte")
problem = find_paragraph(doc, "El planteamiento del problema")
marco = find_paragraph(doc, "MARCO TEÓRICO")
move_range_after(state_start, problem, marco)

replace_text(problem, "1.2 Planteamiento del problema")
problem.style = "Heading 2"
remove_numbering(problem)
context = find_paragraph(doc, "Contexto del caso de estudio")
replace_text(context, "1.2.1 Contexto del caso de estudio")
remove_numbering(context)
question_heading = find_paragraph(doc, "Pregunta de Investigación")
replace_text(question_heading, "1.2.2 Pregunta de investigación")
remove_numbering(question_heading)

question = doc.paragraphs[paragraph_index(doc, question_heading) + 1]
new_question = (
    "¿En qué medida un modelo híbrido de métodos estadísticos y aprendizaje automático reduce, frente a modelos de referencia y mediante "
    "validación temporal, el error fuera de muestra del importe semanal de compras elegibles y de su distribución porcentual por insumo en los "
    "registros de Cup&Cake, microempresa de repostería creativa de Tizayuca, Hidalgo, disponibles del 5 de diciembre de 2016 al 11 de mayo de 2026 "
    "y evaluados del 24 de marzo al 28 de julio de 2025?"
)
replace_text(question, new_question)

objective_heading = find_paragraph(doc, "Objetivo")
hypothesis_heading = find_paragraph(doc, "Hipótesis")
justification_heading = find_paragraph(doc, "Justificación")

# Mueve hipótesis antes de objetivos.
hypothesis_end = marco
move_range_after(hypothesis_heading, hypothesis_end, question)
replace_text(hypothesis_heading, "1.3 Hipótesis y criterio de evaluación")
hypothesis_heading.style = "Heading 2"
remove_numbering(hypothesis_heading)

replace_text(objective_heading, "1.4 Objetivos")
objective_heading.style = "Heading 2"
remove_numbering(objective_heading)
objective_general_label = find_paragraph(doc, "Objetivo general")
replace_text(objective_general_label, "1.4.1 Objetivo general")
objective_general_label.style = "Heading 3"
remove_numbering(objective_general_label)
objective_general = doc.paragraphs[paragraph_index(doc, objective_general_label) + 1]
replace_text(
    objective_general,
    "Determinar en qué medida un modelo híbrido de métodos estadísticos y aprendizaje automático reduce, frente a modelos de referencia y mediante "
    "validación temporal, el error fuera de muestra del importe semanal de compras elegibles y de su distribución porcentual por insumo en los "
    "registros de Cup&Cake, microempresa de repostería creativa de Tizayuca, Hidalgo, disponibles del 5 de diciembre de 2016 al 11 de mayo de 2026 "
    "y evaluados del 24 de marzo al 28 de julio de 2025.",
)
objectives_specific_label = find_paragraph(doc, "Objetivos específicos")
replace_text(objectives_specific_label, "1.4.2 Objetivos específicos")
objectives_specific_label.style = "Heading 3"
remove_numbering(objectives_specific_label)

replace_text(justification_heading, "1.5 Justificación")
remove_numbering(justification_heading)

# Apartados institucionales que hacen explícitos alcance, limitaciones y definiciones.
last_justification = find_paragraph(doc, "La investigación se justifica por la necesidad", exact=False)
scope_h = insert_after(last_justification, "1.6 Alcance", "Heading 2")
remove_numbering(scope_h)
scope_p = insert_after(
    scope_h,
    "El estudio se limita al caso longitudinal-retrospectivo de Cup&Cake. Pronostica el importe monetario semanal de compras elegibles y distribuye "
    "el presupuesto entre categorías registradas para horizontes de una a cuatro semanas. Evalúa precisión fuera de muestra; no estima cantidades "
    "físicas óptimas, inventarios de seguridad, causalidad, rentabilidad ni una política automática de compras.",
    "Normal",
)
limitations_h = insert_after(scope_p, "1.7 Limitaciones", "Heading 2")
remove_numbering(limitations_h)
limitations_p = insert_after(
    limitations_h,
    "La evidencia está condicionada por seis orígenes monetarios por horizonte, entre tres y cinco observaciones composicionales válidas, semanas sin "
    "captura y ausencia de certificación de exhaustividad de los registros. El pronóstico operativo conservado corresponde a un origen histórico de "
    "agosto de 2025. Estas condiciones restringen la generalización y justifican el carácter exploratorio de la evaluación.",
    "Normal",
)
definitions_h = insert_after(limitations_p, "1.8 Definiciones operativas", "Heading 2")
remove_numbering(definitions_h)
insert_after(
    definitions_h,
    "En esta tesis, compras elegibles designa el importe registrado después de aplicar las exclusiones presupuestarias documentadas; modelo híbrido "
    "designa una combinación o corrección residual de componentes estadísticos y de aprendizaje automático; y precisión fuera de muestra designa el "
    "error calculado sobre semanas posteriores al origen, no utilizadas para ajustar ni seleccionar el modelo.",
    "Normal",
)

# Capítulo II y numeración del estado del arte y marco teórico.
replace_text(marco, "CAPÍTULO II MARCO TEÓRICO")
remove_numbering(marco)

chapter2_h2 = [
    "Estado del Arte",
    "Microempresa y planeación del presupuesto de abastecimiento",
    "Compras monetarias y contexto de la repostería",
    "Series temporales semanales y calidad de los registros",
    "Métodos estadísticos y de aprendizaje automático",
    "Fundamentos de integración del modelo híbrido",
    "Distribución porcentual y coherencia de las asignaciones",
    "Datos limitados e ingeniería de características",
    "Evaluación temporal y métricas de precisión",
    "Inteligencia de negocios y comunicación del pronóstico",
    "Síntesis de fundamentos y relación con la hipótesis",
]
for number, old in enumerate(chapter2_h2, 1):
    p = find_paragraph(doc, old)
    replace_text(p, f"2.{number} {old if old != 'Estado del Arte' else 'Revisión de evidencia empírica y estado del arte'}")
    p.style = "Heading 2"
    remove_numbering(p)

# Numera los Heading 3/4 del Capítulo II en el orden actual.
chapter2_start = find_paragraph(doc, "CAPÍTULO II MARCO TEÓRICO")
chapter3_original = find_paragraph(doc, "METODOLOGÍA")
current_h2 = None
sub = 0
subsub = 0
for p in doc.paragraphs:
    if p._p is chapter2_start._p:
        active = True
        continue
    if p._p is chapter3_original._p:
        break
    if not locals().get("active", False):
        continue
    style = p.style.name if p.style else ""
    if style == "Heading 2":
        m = re.match(r"2\.(\d+)", norm(p.text))
        current_h2 = int(m.group(1)) if m else current_h2
        sub = 0
        subsub = 0
    elif style == "Heading 3" and current_h2:
        sub += 1
        subsub = 0
        replace_text(p, f"2.{current_h2}.{sub} {norm(p.text)}")
        remove_numbering(p)
    elif style == "Heading 4" and current_h2:
        subsub += 1
        replace_text(p, f"2.{current_h2}.{sub}.{subsub} {norm(p.text)}")
        remove_numbering(p)

# Reescribe la contribución como aporte técnico delimitado.
contribution = find_paragraph(doc, "La contribución consiste en desarrollar un modelo híbrido", exact=False)
replace_text(
    contribution,
    "La aportación técnica de la investigación no consiste en proponer un algoritmo nuevo. Consiste en diseñar, implementar y evaluar de forma "
    "reproducible un procedimiento híbrido para pronosticar compras elegibles y distribuir el importe entre categorías presupuestarias en una "
    "microempresa con datos limitados, demanda irregular y cobertura incompleta. El procedimiento integra validación temporal, comparación con líneas "
    "base, análisis de sensibilidad, ablación de fuentes y controles de conciliación. Su contribución se delimita al caso, al periodo y a los registros "
    "examinados, sin atribuir superioridad universal a la arquitectura.",
)

# Capítulo III. Desarrollo se integra como implementación del procedimiento.
methodology = chapter3_original
replace_text(methodology, "CAPÍTULO III METODOLOGÍA")
remove_numbering(methodology)

method_h2 = [
    ("Tipo de investigación según el objetivo de esta", "3.1 Tipo y propósito de la investigación"),
    ("Según el alcance evaluativo del estudio", "3.2 Alcance evaluativo del estudio"),
    ("Según el tipo de datos empleados", "3.3 Datos y unidad de análisis"),
    ("Según el grado de manipulación de las variables", "3.4 Diseño no experimental"),
    ("Diseño evaluativo y comparación contra la línea base", "3.5 Diseño evaluativo y líneas base"),
    ("Según el tipo de inferencia", "3.6 Tipo de inferencia"),
    ("Temporalidad de la investigación", "3.7 Delimitación temporal"),
    ("Procedimiento metodológico", "3.8 Procedimiento metodológico"),
]
for old, new in method_h2:
    p = find_paragraph(doc, old)
    replace_text(p, new)
    remove_numbering(p)

method_h3 = [
    "Recolección y estructuración de los conjuntos de datos (Datasets)",
    "Auditoría de cobertura y construcción del panel semanal",
    "Predictores de compras, ventas y calendario",
    "Modelos de referencia y arquitectura híbrida",
    "Validación temporal rolling window",
    "Métricas y criterios de comparación",
    "Comunicación del pronóstico para la planeación presupuestaria",
    "Consideraciones éticas y tratamiento de datos",
]
for number, old in enumerate(method_h3, 1):
    p = find_paragraph(doc, old)
    replace_text(p, f"3.8.{number} {old}")
    remove_numbering(p)

development = find_paragraph(doc, "DESARROLLO")
replace_text(development, "3.9 Implementación y desarrollo del procedimiento")
development.style = "Heading 2"
remove_numbering(development)
development_h2 = [
    "Contexto de los registros y objetivo operativo",
    "Preparación de fuentes y cobertura semanal",
    "Construcción del panel y caracterización exploratoria",
    "Integración de ventas y variables calendario",
    "Definición de categorías y distribución presupuestaria",
    "Modelos y validación temporal",
    "Asignación del presupuesto y controles de calidad",
    "Alcance y trazabilidad de la corrida",
]
for number, old in enumerate(development_h2, 1):
    p = find_paragraph(doc, old)
    replace_text(p, f"3.9.{number} {old}")
    p.style = "Heading 3"
    remove_numbering(p)

# Justificación de configuraciones e hiperparámetros.
architecture = find_paragraph(doc, "3.8.4 Modelos de referencia y arquitectura híbrida")
arch_paras = doc.paragraphs
architecture_index = paragraph_index(doc, architecture)
selection_anchor = arch_paras[architecture_index + 7]
selection_explanation = insert_after(
    selection_anchor,
    "La comparación de configuraciones se realizó exclusivamente en cinco orígenes de validación interna. El criterio monetario fue el RMSE medio por "
    "horizonte; el bloque final permaneció aislado hasta terminar la selección. Los pesos híbridos candidatos fueron 0.25, 0.50 y 0.75, y los parámetros "
    "de decaimiento de la composición fueron 0.10, 0.30 y 0.60. Los modelos de aprendizaje automático conservaron semilla 42 y configuraciones acotadas: "
    "Random Forest con 100 árboles, profundidad máxima 4 y mínimo 3 observaciones por hoja; HistGradientBoosting con 100 iteraciones, 7 hojas máximas y "
    "mínimo 5 observaciones por hoja; Ridge con penalización 10; y MLP con una capa oculta de 4 unidades, activación tanh y optimizador LBFGS. Esta decisión "
    "priorizó estabilidad y capacidad limitada frente al tamaño reducido de la muestra.",
    "Normal",
)

# Tres tablas nuevas de trazabilidad metodológica.
anchor = insert_table_block(
    selection_explanation,
    2,
    "Datos, particiones y periodos utilizados en la evaluación.",
    ["Componente", "Periodo o tamaño", "Uso y restricción"],
    [
        ["Panel disponible", "5 dic 2016–11 may 2026", "Cobertura general; no implica observación exhaustiva de todas las semanas."],
        ["Ventana de entrenamiento", "52 semanas calendario", "Sólo información disponible antes de cada origen; puede contener menos observaciones válidas."],
        ["Validación interna", "5 orígenes; 18 nov 2024–10 feb 2025", "Selecciona familia, combinación, peso y parámetro composicional por horizonte."],
        ["Evaluación final", "6 orígenes; 24 mar–7 jul 2025", "Bloque fuera de muestra; objetivos hasta el 28 jul 2025."],
        ["Horizontes", "1, 2, 3 y 4 semanas", "Pronósticos directos; h=1 es el contraste principal de H1."],
        ["Avance entre orígenes", "3 semanas", "Configuración única aplicada a la corrida canónica."],
        ["Composición", "3–5 observaciones válidas por horizonte", "Se excluyen semanas sin total positivo o sin composición observable."],
    ],
    [1.35, 1.8, 3.0],
    "La Tabla 2 concentra los datos de entrenamiento, selección y evaluación que delimitan qué observaciones intervinieron en cada decisión.",
    "Nota. Los periodos corresponden a la corrida canónica run_20261013T193604Z_157556. Fuente: elaboración propia con base en la configuración y el manifiesto de ejecución.",
    2402,
)

anchor = insert_table_block(
    anchor,
    3,
    "Familias de métodos, configuración y función en la corrida canónica.",
    ["Familia", "Configuración o candidatos", "Función"],
    [
        ["Referencias monetarias", "Último valor, promedio móvil de 4 semanas y referencia estacional cuando existe historia", "Determinar si la complejidad aporta respecto de alternativas sencillas."],
        ["Estadísticos", "SES, Holt amortiguado, SARIMAX AR(1) y ARIMA(1,1,1)/estacional simplificado", "Capturar nivel, tendencia o dependencia lineal."],
        ["Aprendizaje automático", "Ridge, Random Forest, HistGradientBoosting y MLP pequeña", "Representar relaciones regularizadas o no lineales con capacidad acotada."],
        ["Híbridos", "Pesos 0.25, 0.50 y 0.75; correctores residuales RF y MLP sobre ARIMA", "Combinar información complementaria; selección interna independiente por horizonte."],
        ["Composición", "Promedio histórico frente a EWM con alfa 0.10, 0.30 o 0.60", "Distribuir el total elegible y comparar contra una referencia reproducible."],
        ["No ejecutados en la corrida final", "RNN, LSTM, Prophet y modelos log-ratio", "Antecedentes o trabajo futuro; no sustentan los resultados reportados."],
    ],
    [1.4, 2.55, 2.2],
    "La Tabla 3 distingue los métodos realmente ejecutados de los que sólo se revisaron como antecedentes o alternativas futuras.",
    "Nota. La configuración ganadora se eligió por RMSE interno para cada horizonte, sin observar el bloque final. Fuente: elaboración propia a partir del código y los artefactos de selección.",
    2403,
)

anchor = insert_table_block(
    anchor,
    4,
    "Trazabilidad de las categorías utilizadas en el objetivo monetario y composicional.",
    ["Grupo", "Tratamiento", "Razón analítica"],
    [
        ["Categorías elegibles", "Se incluyen en el importe semanal objetivo", "Representan las compras presupuestarias dentro del alcance."],
        ["COMBUSTIBLE, MUEBLES y HERRAMIENTA", "Se excluyen", "No forman parte del abastecimiento recurrente de insumos modelado."],
        ["IMPUTADO e IMPUTADOS", "Se excluyen", "Etiquetas de auditoría; no representan compras observadas."],
        ["OTROS", "Se excluye", "Categoría presupuestaria fuente fuera del objetivo definido."],
        ["11 insumos principales", "Se presentan individualmente", "Permiten interpretar las categorías de mayor relevancia operativa."],
        ["RESTO_ELEGIBLE", "Agrupa categorías elegibles menores", "Conserva el remanente y permite que la asignación concilie con 100%; no equivale a OTROS."],
    ],
    [1.55, 1.65, 2.95],
    "La Tabla 4 explica el paso desde el catálogo original hasta las categorías presentadas en la distribución final.",
    "Nota. La clasificación se fija antes de evaluar modelos y se aplica de forma idéntica a todos los horizontes. Fuente: elaboración propia con base en el catálogo analítico aprobado.",
    2404,
)

# 3. Capítulo IV: fusiona resultados y discusión.
results = find_paragraph(doc, "RESULTADOS")
replace_text(results, "CAPÍTULO IV RESULTADOS Y DISCUSIÓN")
remove_numbering(results)
result_h2 = [
    "Calidad, cobertura e integración de datos",
    "Análisis de sensibilidad de la imputación",
    "Diseño efectivo de evaluación",
    "Rendimiento comparativo del importe semanal",
    "Distribución porcentual por insumo",
    "Contraste de hipótesis",
    "Alcance de los resultados",
]
for number, old in enumerate(result_h2, 1):
    p = find_paragraph(doc, old)
    new = "Evaluación exploratoria de la hipótesis" if old == "Contraste de hipótesis" else old
    replace_text(p, f"4.{number} {new}")
    remove_numbering(p)

discussion = find_paragraph(doc, "DISCUSIÓN")
replace_text(discussion, "4.8 Discusión de resultados")
discussion.style = "Heading 2"
remove_numbering(discussion)
discussion_h2 = [
    "Interpretación de los hallazgos y contraste con antecedentes",
    "Demanda intermitente y selección de métricas",
    "Implicaciones operativas y financieras",
    "Validez y limitaciones",
]
for number, old in enumerate(discussion_h2, 1):
    p = find_paragraph(doc, old)
    replace_text(p, f"4.8.{number} {old}")
    p.style = "Heading 3"
    remove_numbering(p)

# Corrige de manera explícita la interpretación del signo y los parámetros inferenciales.
h1 = find_paragraph(doc, "La evaluación de H1 fue clasificada", exact=False)
replace_text(
    h1,
    "La evaluación de H1 fue retrospectiva y exploratoria. En h=1 se obtuvieron seis pares monetarios y tres pares composicionales. La diferencia "
    "monetaria se definió como pérdida cuadrática del último valor menos pérdida del híbrido; su media de 161,798.64 MXN² fue positiva y favoreció "
    "descriptivamente al híbrido. La diferencia composicional se definió como MAE histórico menos MAE EWM; su media de −1.39 puntos porcentuales fue "
    "negativa y favoreció al promedio histórico. El protocolo exigía un mínimo de 12 pares para inferencia, 2,000 remuestras de bootstrap circular, "
    "bloques de cuatro semanas e intervalos simultáneos Bonferroni para los dos componentes de H1. Como no se alcanzó el mínimo, no se calcularon "
    "intervalos y H1 no quedó confirmada.",
)
for p in doc.paragraphs:
    text = norm(p.text)
    if text == "Las diferencias negativas favorecen al método propuesto. La inferencia se limita a los orígenes y horizontes observados y debe leerse junto con la evidencia monetaria.":
        replace_text(p, "En el contraste composicional, las diferencias negativas favorecen a la referencia histórica. La inferencia se limita a los orígenes y horizontes observados y debe leerse junto con la evidencia monetaria.")
    elif text == "Las diferencias negativas favorecen al modelo propuesto; la conclusión debe considerar su dispersión y el tamaño efectivo de la evaluación, no sólo el signo del promedio.":
        replace_text(p, "En el contraste monetario, los valores positivos favorecen al híbrido frente al último valor. La conclusión considera también la dispersión y el tamaño efectivo de la evaluación, no sólo el signo del promedio.")

ablation_discussion = find_paragraph(doc, "Las ventas se incorporaron como información histórica disponible", exact=False)
replace_text(
    ablation_discussion,
    "Las ventas y el calendario se incorporaron como información disponible al origen. El análisis de ablación comparó configuraciones que retiraron "
    "de forma controlada ventas, calendario o componentes del modelo bajo las mismas particiones. Sus diferencias describen sensibilidad dentro de "
    "esta corrida; por el número reducido de orígenes, no permiten atribuir causalmente una reducción estable del error a una fuente aislada.",
)

# 4. Capítulo V.
conclusions = find_paragraph(doc, "CONCLUSIONES")
replace_text(conclusions, "CAPÍTULO V CONCLUSIONES Y RECOMENDACIONES")
remove_numbering(conclusions)
conclusions_intro = insert_after(conclusions, "5.1 Conclusiones", "Heading 2")
remove_numbering(conclusions_intro)
recommendations = find_paragraph(doc, "Recomendaciones")
replace_text(recommendations, "5.2 Recomendaciones")
remove_numbering(recommendations)
future = find_paragraph(doc, "Trabajo futuro")
replace_text(future, "5.3 Trabajo futuro")
remove_numbering(future)

# 5. Renumera títulos de tablas existentes 2–6 como 5–9 y actualiza marcadores/REF.
body = doc._element.body
for old, new in ((6, 9), (5, 8), (4, 7), (3, 6), (2, 5)):
    old_name = f"_RefTabla{old}"
    new_name = f"_RefTabla{new}"
    for bookmark in body.xpath(".//w:bookmarkStart"):
        if bookmark.get(qn("w:name")) == old_name:
            bookmark.set(qn("w:name"), new_name)
    for field in body.xpath(".//w:fldSimple"):
        instruction = field.get(qn("w:instr"), "")
        if old_name in instruction:
            field.set(qn("w:instr"), instruction.replace(old_name, new_name))
            texts = list(field.iter(qn("w:t")))
            for text_node in texts:
                if text_node.text == str(old):
                    text_node.text = str(new)

for p in doc.paragraphs:
    names = [node.get(qn("w:name"), "") for node in p._p.xpath(".//w:bookmarkStart")]
    for number in range(5, 10):
        if f"_RefTabla{number}" in names:
            for field in p._p.xpath(".//w:fldSimple"):
                if "SEQ Tabla" in field.get(qn("w:instr"), ""):
                    for text_node in field.iter(qn("w:t")):
                        text_node.text = str(number)

existing_table_captions = [p for p in doc.paragraphs if p.style and p.style.name == "Caption" and norm(p.text).startswith("Tabla")]
for p in existing_table_captions:
    plain = norm(p.text)
    for old, new in ((6, 9), (5, 8), (4, 7), (3, 6), (2, 5)):
        if plain.startswith(f"Tabla {old}."):
            replace_text(p, plain.replace(f"Tabla {old}.", f"Tabla {new}.", 1))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break

# Centra todos los títulos de tablas y figuras del cuerpo.
for p in doc.paragraphs:
    if p.style and p.style.name == "Caption":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 6. Paginación institucional por secciones.
sections = list(doc.sections)
for idx, section in enumerate(sections):
    sectpr = section._sectPr
    pgnum = sectpr.find(qn("w:pgNumType"))
    if pgnum is None:
        pgnum = OxmlElement("w:pgNumType")
        sectpr.append(pgnum)
    if idx == 1:
        pgnum.set(qn("w:fmt"), "lowerRoman")
        pgnum.set(qn("w:start"), "1")
    elif idx == 2:
        pgnum.set(qn("w:fmt"), "decimal")
        pgnum.set(qn("w:start"), "1")

# Evita que Word intente actualizar automáticamente campos frágiles al abrir.
settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is not None:
    update.set(qn("w:val"), "false")

doc.core_properties.title = "Modelo híbrido para el pronóstico y la distribución del presupuesto semanal de abastecimiento"
doc.core_properties.comments = "Revisión 64: correcciones del Director y comité, estructura institucional de cinco capítulos y fortalecimiento metodológico."
doc.save(OUT)
print(OUT)
