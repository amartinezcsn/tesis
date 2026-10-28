from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "documentacion" / "TESIS_OCT2026_Rev65_(ZUJ)_Times_New_Roman.docx"
OUT = ROOT / "documentacion" / "TESIS_OCT2026_Rev66_(ZUJ)_Rolling_Window.docx"
FIG = ROOT / "documentacion" / "figuras_rev66" / "Figura_8_rolling_window_52_semanas.png"


REPLACEMENTS = {
    "avance de tres semanas": "avance de una semana",
    "a three-week step": "a one-week step",
    "avance cada tres semanas": "avance cada semana",
    "desplaza el origen cada tres semanas": "desplaza la ventana una semana en cada iteración",
}


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_font(run, name="Times New Roman", size=None):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def replace_text(paragraph, old, new):
    if old not in paragraph.text:
        return False
    full = paragraph.text.replace(old, new)
    clear_paragraph(paragraph)
    set_font(paragraph.add_run(full))
    return True


def set_paragraph(paragraph, text, font="Times New Roman"):
    clear_paragraph(paragraph)
    set_font(paragraph.add_run(text), font)


def replace_picture(paragraph, image_path):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ppr = paragraph._p.get_or_add_pPr()
    keep_next = ppr.find(qn("w:keepNext"))
    if keep_next is None:
        keep_next = OxmlElement("w:keepNext")
        ppr.append(keep_next)
    run = paragraph.add_run()
    set_font(run)
    run.add_picture(str(image_path), width=Inches(6.05))


def all_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def main():
    assert SRC.exists(), SRC
    assert FIG.exists(), FIG
    doc = Document(SRC)

    counts = {key: 0 for key in REPLACEMENTS}
    for paragraph in all_paragraphs(doc):
        for old, new in REPLACEMENTS.items():
            if replace_text(paragraph, old, new):
                counts[old] += 1

    # Matiza el diseño y los tamaños muestrales con la evidencia de particiones.
    exact = {
        "La configuración aplicada utiliza 52 semanas de entrenamiento, avance de una semana y horizontes de una a cuatro semanas. La validación interna utiliza cinco orígenes para elegir componentes; la disponibilidad de historia y cobertura determina qué folds pueden calcularse. Los objetivos de la última partición interna preceden a la evaluación final para evitar contaminación temporal.":
        "La configuración aplicada utiliza una rolling window de 52 semanas y un desplazamiento semanal. En cada iteración sale la semana más antigua y entra la semana disponible más reciente, por lo que el ancho del entrenamiento permanece constante. La validación interna utiliza cinco ventanas para elegir componentes; los objetivos de la última partición interna preceden a la evaluación final para evitar contaminación temporal.",
        "La evaluación usa rolling window con 52 semanas de entrenamiento, avance cada semana y horizontes directos de una a cuatro semanas. Cinco orígenes internos se utilizaron para seleccionar componentes por horizonte. La corrida final produjo seis folds observados por horizonte; los modelos estacionales de 52 semanas tuvieron menos folds comparables por disponibilidad. Las exclusiones y tamaños muestrales se reportan sin rellenar artificialmente los objetivos faltantes.":
        "La evaluación usa rolling window con una ventana fija de 52 semanas, desplazamiento semanal y horizontes directos de una a cuatro semanas. Se generaron cinco ventanas internas para seleccionar componentes por horizonte y 17 ventanas candidatas para la evaluación final. Las métricas incorporan únicamente objetivos observados, lo que produce seis folds válidos por horizonte; los modelos estacionales de 52 semanas tuvieron menos comparaciones por disponibilidad. Las exclusiones y tamaños muestrales se reportan sin rellenar artificialmente los objetivos faltantes.",
        "La combinación ponderada, la corrección residual ARIMA+MLP y la generación de orígenes rolling se formalizan en las ecuaciones (37) a (39).":
        "La combinación ponderada, la corrección residual ARIMA+MLP y la secuencia de ventanas deslizantes se formalizan en las ecuaciones (37) a (39).",
        "La Ecuación 39 formaliza los orígenes, la ventana de entrenamiento y la semana objetivo. Esta notación garantiza que la evaluación respete el orden temporal.":
        "La Ecuación 39 formaliza una secuencia semanal de ventanas. Para cada origen, el intervalo de entrenamiento conserva 52 semanas y el objetivo se ubica según el horizonte h; así, la evaluación respeta el orden temporal sin ampliar acumulativamente el historial.",
        "La evaluación reproduce el momento en que se emitiría el presupuesto y mantiene el orden temporal de la información. Como se muestra en la Figura , la evidencia se presenta de manera independiente.":
        "La Figura 8 demuestra la mecánica aplicada: las barras de entrenamiento conservan el mismo ancho de 52 semanas, cada ventana se desplaza una semana y los objetivos h=1 a h=4 permanecen fuera del intervalo de entrenamiento.",
        "Cada origen utiliza únicamente el historial disponible y reserva las semanas posteriores para medir error, lo que evita una comparación optimista por fuga temporal.":
        "Cada ventana utiliza únicamente las 52 semanas anteriores a su origen. El origen funciona como límite temporal —no como nombre de la técnica— y las semanas posteriores se reservan para medir el error, lo que evita fuga de información y comparaciones optimistas.",
    }
    changed_exact = 0
    for paragraph in doc.paragraphs:
        if paragraph.text in exact:
            set_paragraph(paragraph, exact[paragraph.text])
            changed_exact += 1
    assert changed_exact == len(exact), (changed_exact, len(exact))

    # Tabla de configuración metodológica.
    table = doc.tables[30]
    table.cell(3, 1).text = "5 ventanas; 20 ene–17 feb 2025"
    table.cell(3, 2).text = "Seleccionan familia, combinación, peso y parámetro composicional por horizonte."
    table.cell(4, 1).text = "17 ventanas candidatas; 17 mar–7 jul 2025"
    table.cell(4, 2).text = "Bloque fuera de muestra; seis folds con objetivo observado por horizonte y objetivos hasta el 28 jul 2025."
    table.cell(6, 1).text = "1 semana"
    table.cell(6, 2).text = "Desplazamiento semanal aplicado a la ventana fija de 52 semanas."

    # Matriz de sustento teórico.
    table0 = doc.tables[0]
    table0.cell(7, 7).text = (
        "Sustenta la validación temporal mediante ventanas deslizantes y el uso cauteloso "
        "de métricas; small data corresponde a la condición de esta tesis."
    )

    # Ecuación 39: una nueva ventana en cada semana, no cada tres semanas.
    eq_cell = doc.tables[41].cell(0, 1)
    eq_p = eq_cell.paragraphs[0]
    set_paragraph(eq_p, "oₖ = o₀ + k;   Tₒ = {t : o−52 ≤ t < o};   objetivo = Yₒ₊ₕ₋₁", "Cambria Math")
    eq_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Figura 8 y pie de figura.
    caption = next(p for p in doc.paragraphs if p.style and p.style.name == "Caption" and
                   p.text.startswith("Figura ") and
                   "Esquema de validación temporal con orígenes sucesivos" in p.text)
    doc_paragraphs = list(doc.paragraphs)
    cap_idx = next(i for i, p in enumerate(doc_paragraphs) if p._p is caption._p)
    image_p = next(p for p in reversed(doc_paragraphs[:cap_idx]) if p._p.xpath(".//w:drawing|.//w:pict"))
    replace_picture(image_p, FIG)
    set_paragraph(
        caption,
        "Figura 8. Esquema de validación rolling window con ventanas fijas de 52 semanas, "
        "desplazamiento semanal y horizontes de una a cuatro semanas. Fuente: elaboración propia "
        "con las particiones de la corrida run_20261013T193604Z_157556."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Actualiza la entrada estática del índice de figuras; la página se corrige tras renderizar.
    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.name.lower().startswith("toc") and \
                "Esquema de validación temporal con orígenes sucesivos" in paragraph.text:
            page = paragraph.text.rsplit("\t", 1)[-1]
            set_paragraph(paragraph,
                "Figura 8. Esquema de validación rolling window con ventanas fijas de 52 semanas, "
                f"desplazamiento semanal y horizontes de una a cuatro semanas.\t{page}")

    # Homologa el texto nuevo a Times New Roman, manteniendo Cambria Math sólo en la ecuación.
    for paragraph in all_paragraphs(doc):
        is_equation = paragraph is eq_p
        for run in paragraph.runs:
            set_font(run, "Cambria Math" if is_equation else "Times New Roman")

    # python-docx recrea objetos Paragraph al recorrer tablas; reafirma la fuente matemática.
    for run in eq_p.runs:
        set_font(run, "Cambria Math")

    doc.core_properties.title = "Tesis doctoral — versión 66 — validación Rolling Window"
    doc.save(OUT)
    print(OUT)
    print(counts)


if __name__ == "__main__":
    main()
