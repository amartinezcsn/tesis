from pathlib import Path
import re
import sys
import unicodedata

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


DOCX = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"C:\Python\tesis\documentacion\TESIS_OCT2026_Rev66_(ZUJ)_Rolling_Window.docx")
PDF = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(r"C:\Python\tesis\tmp\rev66_render\TESIS_OCT2026_Rev66_(ZUJ)_Rolling_Window.pdf")


def key(text):
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(re.findall(r"[a-z0-9]+", text.lower()))


def roman(number):
    values = (
        (1000, "m"), (900, "cm"), (500, "d"), (400, "cd"),
        (100, "c"), (90, "xc"), (50, "l"), (40, "xl"),
        (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"),
    )
    result = []
    for value, symbol in values:
        while number >= value:
            result.append(symbol)
            number -= value
    return "".join(result)


reader = PdfReader(PDF)
page_texts = [page.extract_text() or "" for page in reader.pages]
page_keys = [key(text) for text in page_texts]
body_query = key("CAPÍTULO I INTRODUCCIÓN")
body_page = max(i + 1 for i, content in enumerate(page_keys) if body_query in content)


def locate(title):
    query = key(title)
    if query in {"glosario", "referencias"}:
        for i, content in enumerate(page_texts):
            if query in {key(line) for line in content.splitlines()}:
                return i + 1
    candidates = [query]
    words = query.split()
    if len(words) > 12:
        candidates.append(" ".join(words[:12]))
    candidates.append(re.sub(r"^(tabla|figura) \d+ ", "", query))
    for candidate in candidates:
        if len(candidate) < 4:
            continue
        for i, content in enumerate(page_keys[body_page - 3 :], body_page - 3):
            if candidate in content:
                return i + 1
    raise ValueError(f"No se encontró en PDF: {title}")


doc = Document(DOCX)
updated = 0
for paragraph in doc.paragraphs:
    style = paragraph.style.name.lower() if paragraph.style else ""
    if not style.startswith("toc ") or "\t" not in paragraph.text:
        continue
    title = paragraph.text.rsplit("\t", 1)[0].strip()
    physical = locate(title)
    label = roman(physical - 1) if physical < body_page else str(physical - body_page + 1)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.add_run(f"{title}\t{label}")
    updated += 1

assert updated == 141, updated
doc.save(DOCX)
print(f"Páginas físicas: {len(reader.pages)}; inicio del cuerpo: {body_page}; entradas: {updated}")
