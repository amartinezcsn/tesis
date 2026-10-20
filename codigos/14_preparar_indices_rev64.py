from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOCX = Path(r"C:\Python\tesis\documentacion\TESIS_OCT2026_Rev64_(ZUJ)_Correcciones_Director.docx")


def norm(text):
    return " ".join(text.split())


def visible_text(paragraph):
    return norm("".join(node.text or "" for node in paragraph._p.iter(qn("w:t"))))


def insert_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p.getparent().remove(new_para._p)
    new_p.getparent().replace(new_p, new_para._p)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def find(doc, text):
    for p in doc.paragraphs:
        if norm(p.text) == text:
            return p
    raise ValueError(text)


def clear_between(start, end, styles):
    node = start._p.getnext()
    while node is not None and node is not end._p:
        nxt = node.getnext()
        pstyle = node.find(qn("w:pPr") + "/" + qn("w:pStyle"))
        style_id = pstyle.get(qn("w:val"), "") if pstyle is not None else ""
        if style_id.lower() in styles or style_id.lower().replace("tdc", "toc") in styles:
            node.getparent().remove(node)
        node = nxt


doc = Document(DOCX)

# Captions del cuerpo antes de modificar índices.
tables = []
figures = []
for p in doc.paragraphs:
    if p.style and p.style.name == "Caption":
        text = visible_text(p)
        if text.startswith("Tabla "):
            tables.append(text.split(" Fuente:")[0])
        elif text.startswith("Figura "):
            figures.append(text.split(" Fuente:")[0])

assert len(tables) == 9, tables
assert len(figures) == 21, figures
tables.sort(key=lambda text: int(text.split(".", 1)[0].split()[1]))

# Índice general estático con páginas provisionales.
toc_heading = find(doc, "ÍNDICE")
tables_heading = find(doc, "ÍNDICE DE TABLAS")
clear_between(toc_heading, tables_heading, {"toc1", "toc2", "toc3", "toc4"})

entries = []
started = False
for p in doc.paragraphs:
    text = norm(p.text)
    if text == "RESUMEN":
        started = True
    if not started or not text:
        continue
    style = p.style.name if p.style else ""
    if style in {"Heading 1", "Heading 2", "Heading 3", "Heading 4"}:
        if text in {"RESUMEN", "ABSTRACT"}:
            level = 1
        else:
            level = int(style.rsplit(" ", 1)[1])
        entries.append((level, text))

anchor = toc_heading
for level, title in entries:
    anchor = insert_after(anchor, f"{title}\tPÁG", f"toc {level}")

# Índice de tablas.
figures_heading = find(doc, "ÍNDICE DE FIGURAS")
clear_between(tables_heading, figures_heading, {"toc1", "toc2", "toc3", "toc4"})
anchor = tables_heading
for title in tables:
    anchor = insert_after(anchor, f"{title}\tPÁG", "toc 1")

# Índice de figuras hasta el primer encabezado posterior (RESUMEN).
summary_heading = find(doc, "RESUMEN")
clear_between(figures_heading, summary_heading, {"toc1", "toc2", "toc3", "toc4"})
anchor = figures_heading
for title in figures:
    anchor = insert_after(anchor, f"{title}\tPÁG", "toc 1")

doc.save(DOCX)
print(f"TOC: {len(entries)}; tablas: {len(tables)}; figuras: {len(figures)}")
