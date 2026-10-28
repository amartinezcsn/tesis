from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "documentacion" / "TESIS_OCT2026_Rev66_(ZUJ)_Rolling_Window.docx"
OUT = ROOT / "documentacion" / "TESIS_OCT2026_Rev67_(ZUJ)_Glosario_Complementado.docx"


ENTRIES = [
    ("Ablación", "Comparación controlada en la que se retira una fuente, grupo de variables o componente del modelo para estimar su contribución al desempeño predictivo."),
    ("Aprendizaje automático", "Conjunto de métodos computacionales que ajustan patrones a partir de datos para realizar predicciones o clasificaciones sin programar explícitamente cada regla de decisión."),
    ("ARIMA", "Modelo autorregresivo integrado de medias móviles que representa la dependencia temporal mediante términos autorregresivos, diferenciación y errores rezagados."),
    ("Autocorrelación", "Asociación lineal entre una serie temporal y sus propios valores observados en rezagos anteriores."),
    ("Bootstrap", "Procedimiento de remuestreo utilizado para aproximar la variabilidad de una estimación. En esta tesis no produjo intervalos para H1 porque el número de pares válidos no alcanzó el mínimo establecido."),
    ("Cadena de suministro", "Conjunto de procesos relacionados con el abastecimiento, la transformación y la entrega de bienes, desde los proveedores hasta el cliente final."),
    ("Categoría analítica", "Agrupación normalizada de insumos utilizada para integrar registros, calcular participaciones y presentar la asignación presupuestaria."),
    ("Conciliación aritmética", "Condición por la que la suma de las asignaciones por categoría coincide con el importe total pronosticado, salvo diferencias menores de redondeo."),
    ("Datos composicionales", "Vectores de participaciones no negativas cuya suma está restringida a un total fijo; en esta tesis, 100 % del importe elegible semanal."),
    ("Demanda intermitente", "Patrón temporal con numerosos periodos sin valor positivo y apariciones irregulares de importes distintos de cero."),
    ("Estacionariedad", "Propiedad aproximada por la que características estadísticas de una serie, como nivel, variabilidad o dependencia, permanecen estables en el tiempo."),
    ("Fold temporal", "Unidad de evaluación formada por una ventana de entrenamiento, un origen y uno o más objetivos posteriores. Sólo los folds con objetivo observado intervienen en las métricas."),
    ("Fuga de información", "Uso directo o indirecto de datos futuros o del bloque de evaluación durante el entrenamiento o la selección del modelo, lo que genera estimaciones optimistas del desempeño."),
    ("HistGradientBoosting", "Método de potenciación de gradiente basado en histogramas que combina secuencialmente árboles de decisión para reducir los errores residuales."),
    ("Horizonte de pronóstico", "Número de semanas entre el origen y la fecha objetivo. La tesis evalúa horizontes directos de una a cuatro semanas, denotados h=1, h=2, h=3 y h=4."),
    ("Importe elegible", "Suma monetaria semanal de las compras que cumplen los criterios de inclusión y pertenecen a las categorías consideradas en el objetivo del estudio."),
    ("Imputación", "Sustitución explícita de valores ausentes mediante una regla de estimación. Las semanas estimadas en productos de auditoría no se incorporaron al ajuste ni a la evaluación de la corrida canónica."),
    ("Inteligencia artificial (IA)", "Campo de la informática que desarrolla sistemas capaces de ejecutar tareas asociadas con percepción, razonamiento, aprendizaje, planificación o predicción."),
    ("Línea base", "Método sencillo y reproducible utilizado como referencia para determinar si un modelo más complejo aporta una mejora predictiva verificable."),
    ("MAE", "Error absoluto medio; promedio de las magnitudes absolutas de los errores. Se expresa en la misma unidad que la variable pronosticada."),
    ("MASE", "Error absoluto medio escalado; cociente entre el MAE del modelo y un error de referencia calculado con el historial de entrenamiento. Permite comparar escalas, aunque su lectura depende del denominador disponible."),
    ("MLP", "Perceptrón multicapa; red neuronal de alimentación hacia adelante formada por capas de unidades conectadas. En la tesis se emplea una arquitectura pequeña y no recurrente."),
    ("Modelo híbrido", "Procedimiento que integra pronósticos estadísticos y de aprendizaje automático mediante combinación ponderada o corrección de residuos."),
    ("Origen del pronóstico", "Fecha que separa la información disponible para entrenar de las semanas futuras que se pronostican. Es el límite de cada fold, no el nombre de la técnica de validación."),
    ("OTROS", "Etiqueta reservada para categorías explícitamente clasificadas como OTROS en los registros. No es equivalente a RESTO_ELEGIBLE."),
    ("Panel semanal", "Conjunto de datos organizado por semanas calendario que integra la variable objetivo, participaciones, ventas y predictores disponibles para cada periodo."),
    ("Participación porcentual", "Proporción del importe elegible atribuida a una categoría analítica. Las participaciones de una semana válida suman 100 %."),
    ("Pronóstico de demanda", "Estimación de valores futuros de una variable de demanda mediante información disponible hasta una fecha de corte y un método reproducible."),
    ("PYMES", "Pequeñas y medianas empresas; organizaciones de escala limitada en personal y recursos. La empresa del caso se caracteriza específicamente como microempresa."),
    ("Random Forest", "Método de aprendizaje automático que combina numerosos árboles de decisión ajustados con muestras y subconjuntos de variables para reducir la variabilidad predictiva."),
    ("Regresión Ridge", "Modelo lineal regularizado que penaliza la magnitud cuadrática de los coeficientes para controlar la complejidad y la inestabilidad por variables correlacionadas."),
    ("Residuo", "Diferencia entre el valor observado y el valor estimado por un modelo. En la alternativa residual, una MLP intenta predecir el error del componente ARIMA."),
    ("RESTO_ELEGIBLE", "Categoría de conciliación que reúne insumos elegibles de menor participación no mostrados individualmente; permanece dentro del total elegible y no representa OTROS."),
    ("RMSE", "Raíz del error cuadrático medio; métrica que eleva al cuadrado los errores antes de promediarlos, por lo que penaliza con mayor intensidad las desviaciones grandes. Se expresa en la unidad original."),
    ("Rolling Window", "Validación temporal con una ventana de entrenamiento de ancho fijo que se desplaza a través del tiempo. En la corrida canónica conserva 52 semanas y avanza una semana en cada iteración."),
    ("Serie temporal", "Secuencia de observaciones ordenadas cronológicamente, registrada en esta tesis con frecuencia semanal."),
    ("Tablero de inteligencia de negocios", "Interfaz que presenta pronósticos, asignaciones, métricas y controles para apoyar la revisión humana de la planeación; no constituye una orden automática de compra."),
    ("Trazabilidad", "Capacidad de relacionar datos, transformaciones, configuraciones, particiones, resultados y artefactos con la corrida que los produjo."),
    ("Variable exógena", "Predictor externo a la variable objetivo que sólo puede utilizarse si su valor está disponible al momento de emitir el pronóstico."),
    ("Ventana de entrenamiento", "Intervalo histórico utilizado para ajustar un modelo en cada fold. En este estudio abarca las 52 semanas anteriores al origen y mantiene un ancho constante."),
]


def set_font(run, name="Times New Roman", size=12):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def format_entry(paragraph, term, definition):
    clear_paragraph(paragraph)
    paragraph.style = "Normal"
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0)
    fmt.first_line_indent = Inches(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.15
    term_run = paragraph.add_run(f"{term}: ")
    term_run.bold = True
    set_font(term_run)
    definition_run = paragraph.add_run(definition)
    set_font(definition_run)


def insert_before(reference_paragraph):
    new_p = OxmlElement("w:p")
    reference_paragraph._p.addprevious(new_p)
    from docx.text.paragraph import Paragraph
    return Paragraph(new_p, reference_paragraph._parent)


def main():
    assert SRC.exists(), SRC
    doc = Document(SRC)
    glossary = next(p for p in doc.paragraphs if p.text.strip() == "GLOSARIO")
    references = next(p for p in doc.paragraphs if p.text.strip() == "REFERENCIAS")
    paragraphs = list(doc.paragraphs)
    start = next(i for i, p in enumerate(paragraphs) if p._p is glossary._p)
    end = next(i for i, p in enumerate(paragraphs) if p._p is references._p)
    old_entries = paragraphs[start + 1:end]
    assert len(old_entries) == 5, len(old_entries)

    for paragraph, (term, definition) in zip(old_entries, ENTRIES[:5]):
        format_entry(paragraph, term, definition)
    for term, definition in ENTRIES[5:]:
        format_entry(insert_before(references), term, definition)

    # El glosario debe iniciar en página nueva y permanecer enlazado al índice general.
    glossary.paragraph_format.page_break_before = True
    doc.core_properties.title = "Tesis doctoral versión 67 glosario complementado"
    doc.save(OUT)
    print(OUT)
    print(f"Entradas anteriores: {len(old_entries)}; entradas revisadas: {len(ENTRIES)}")


if __name__ == "__main__":
    main()
