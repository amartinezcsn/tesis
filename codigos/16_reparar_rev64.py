from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


DOCX = Path(r"C:\Python\tesis\documentacion\TESIS_OCT2026_Rev64_(ZUJ)_Correcciones_Director.docx")


def text_of(element):
    return " ".join("".join(n.text or "" for n in element.iter(qn("w:t"))).split())


def paragraph_style_id(element):
    style = element.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return style.get(qn("w:val"), "") if style is not None else ""


def find_body_index(body, exact_text):
    for i, element in enumerate(list(body)):
        if text_of(element) == exact_text:
            return i
    raise ValueError(exact_text)


def remove_range(body, start, end):
    for element in list(body)[start:end]:
        body.remove(element)


def replace_plain(paragraph, text):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.add_run(text)


doc = Document(DOCX)
body = doc._element.body

# Elimina el índice general automático obsoleto: comienza en el primer TDC
# con instrucciones de campo después del índice estático y termina antes del
# encabezado del índice de tablas.
idx_tables = find_body_index(body, "ÍNDICE DE TABLAS")
stale_toc_start = None
for i, element in enumerate(list(body)[:idx_tables]):
    instruction = "".join(n.text or "" for n in element.iter(qn("w:instrText")))
    if paragraph_style_id(element).upper().startswith("TDC") and instruction:
        stale_toc_start = i
        break
if stale_toc_start is not None:
    remove_range(body, stale_toc_start, idx_tables)

# Conserva sólo las entradas estáticas recién creadas en los índices de tablas
# y figuras; elimina los remanentes del índice automático anterior.
idx_tables = find_body_index(body, "ÍNDICE DE TABLAS")
idx_figures = find_body_index(body, "ÍNDICE DE FIGURAS")
table_entries = [
    i for i in range(idx_tables + 1, idx_figures)
    if paragraph_style_id(list(body)[i]).upper().startswith("TDC")
]
for i in reversed(table_entries[9:]):
    body.remove(list(body)[i])

idx_figures = find_body_index(body, "ÍNDICE DE FIGURAS")
idx_summary = find_body_index(body, "RESUMEN")
figure_entries = [
    i for i in range(idx_figures + 1, idx_summary)
    if paragraph_style_id(list(body)[i]).upper().startswith("TDC")
]
for i in reversed(figure_entries[21:]):
    body.remove(list(body)[i])

# Numera las tablas según su orden real de aparición. Los títulos quedan como
# texto estable para evitar que Word recupere valores almacenados de campos SEQ.
table_captions = [
    p for p in doc.paragraphs
    if p.style and p.style.name == "Caption" and p.text.startswith("Tabla")
]
assert len(table_captions) == 9, len(table_captions)
for number, paragraph in enumerate(table_captions, 1):
    caption = paragraph.text
    caption = re.sub(r"^Tabla\s*(?:\d+)?\s*\.\s*", "", caption)
    replace_plain(paragraph, f"Tabla {number}. {caption}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Corrige las referencias narrativas que aún conservaban la numeración de la
# versión anterior o un campo roto.
reference_replacements = {
    "el desempeño final se resume en la Tabla 3": "el desempeño final se resume en la Tabla 6",
    "Los controles de coherencia se resumen en la Tabla ,": "Los controles de coherencia se resumen en la Tabla 8,",
    "La Tabla 5 resume la igualdad": "La Tabla 9 resume la igualdad",
}
for paragraph in doc.paragraphs:
    original = paragraph.text
    revised = original
    for old, new in reference_replacements.items():
        revised = revised.replace(old, new)
    if revised != original:
        replace_plain(paragraph, revised)

# Evita la expansión del espaciado de palabras observada en el título largo del
# capítulo V y uniforma la presentación de los cinco capítulos.
for paragraph in doc.paragraphs:
    if paragraph.style and paragraph.style.name == "Heading 1" and paragraph.text.startswith("CAPÍTULO"):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.save(DOCX)
print(DOCX)
