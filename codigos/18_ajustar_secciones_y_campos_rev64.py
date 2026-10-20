from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


DOCX = Path(r"C:\Python\tesis\documentacion\TESIS_OCT2026_Rev64_(ZUJ)_Correcciones_Director.docx")

doc = Document(DOCX)
body = doc._element.body
children = list(body)

# La sección decimal comenzaba por error en ABSTRACT. Traslada el salto de
# sección al párrafo inmediatamente anterior al capítulo I, de modo que RESUMEN
# y ABSTRACT permanezcan en romanos y el cuerpo inicie en 1.
chapter = next(p for p in doc.paragraphs if p.text == "CAPÍTULO I INTRODUCCIÓN")
chapter_index = children.index(chapter._p)
source_paragraph = [
    element for element in children[:chapter_index]
    if element.tag == qn("w:p")
    and element.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None
    and any(
        "ABSTRACT" in "".join(n.text or "" for n in later.iter(qn("w:t")))
        for later in children[children.index(element) + 1 : chapter_index]
    )
][-1]
source_ppr = source_paragraph.find(qn("w:pPr"))
sectpr = source_ppr.find(qn("w:sectPr"))
source_ppr.remove(sectpr)

target = children[chapter_index - 1]
if target.tag != qn("w:p"):
    raise RuntimeError("El elemento anterior al capítulo I no es un párrafo")
target_ppr = target.find(qn("w:pPr"))
if target_ppr is None:
    target_ppr = chapter._p.makeelement(qn("w:pPr"))
    target.insert(0, target_ppr)
target_ppr.append(sectpr)

# Elimina el último REF de tabla heredado. La referencia narrativa ya es texto
# explícito y el marcador de la versión anterior provocaba un error al exportar.
for paragraph in doc.paragraphs:
    for field in list(paragraph._p.xpath(".//w:fldSimple")):
        if "_RefTabla" in field.get(qn("w:instr"), ""):
            field.getparent().remove(field)

doc.save(DOCX)
print(DOCX)
