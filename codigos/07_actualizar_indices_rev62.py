from pathlib import Path
import re
from docx import Document
from docx.oxml import OxmlElement


DOCX = Path(r"C:\Python\tesis\documentacion\TESIS_SEP2026_Rev62_(ZUJ)_Final.docx")
FIGURE_PAGES = {
    1: 68, 2: 69, 3: 70, 4: 71, 5: 72, 6: 73, 7: 74, 8: 75,
    9: 77, 10: 78, 11: 80, 12: 81, 13: 84, 14: 90, 15: 92,
    16: 93, 17: 94, 18: 95, 19: 117, 20: 118, 21: 119,
}


def clear_paragraph(p):
    for child in list(p._p):
        if not child.tag.endswith("}pPr"):
            p._p.remove(child)


def insert_after(paragraph, text, style):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    result = paragraph._parent.add_paragraph()
    result._p.getparent().remove(result._p)
    new_p.getparent().replace(new_p, result._p)
    result.style = style
    result.add_run(text)
    return result


doc = Document(DOCX)
paras = doc.paragraphs

# Índice de imágenes
start = next(i for i, p in enumerate(paras) if p.text.strip() == "ÍNDICE DE IMÁGENES")
entries = []
for p in paras[start + 1:]:
    if p.style and p.style.name == "toc 1" and p.text.strip().startswith("Figura"):
        entries.append(p)
    elif entries:
        break
for p in entries:
    match = re.match(r"Figura\s+(\d+)\.\s*(.*?)(?:\tPÁG|\t\d+)?$", p.text.strip())
    if not match:
        continue
    n = int(match.group(1))
    title = match.group(2).rstrip(". ")
    clear_paragraph(p)
    p.add_run(f"Figura {n}. {title}.\t{FIGURE_PAGES[n]}")

# Índice de tablas
for p in doc.paragraphs:
    if p.style and p.style.name == "toc 1" and p.text.strip().startswith("Tabla 6."):
        title = p.text.split("\t", 1)[0].rstrip(". ")
        clear_paragraph(p)
        p.add_run(f"{title}.\t85")

# Índice general: incorpora el anexo nuevo sin actualizar el campo heredado.
paras = doc.paragraphs
general_start = next(i for i, p in enumerate(paras) if p.text.strip() == "ÍNDICE")
anexo_b = next((p for p in paras[general_start + 1:] if p.style and p.style.name == "toc 1" and p.text.strip().startswith("Anexo B")), None)
if anexo_b is not None and not any(p.text.strip().startswith("Anexo C") for p in paras[general_start + 1:start]):
    insert_after(anexo_b, "Anexo C Diagnósticos complementarios y sensibilidad\t117", "toc 1")

doc.save(DOCX)
print(DOCX)
