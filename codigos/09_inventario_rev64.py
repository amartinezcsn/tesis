from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


ROOT = Path(r"C:\Python\tesis")
FILES = [
    ROOT / "documentacion" / "TESIS_SEP2026_Rev62_(ZUJ)_Final.docx",
    ROOT / "documentacion" / "TESIS_OCT2026_Rev63_(ZUJ)_Revision_Comite.docx",
    ROOT / "documentacion" / "OBSERVACIONES 1 y 2.docx",
    ROOT / "documentacion" / "Analisis__tesis_RODRIGO.docx",
    ROOT / "documentacion" / "TESIS_OCT2026_Rev64_(ZUJ)_Correcciones_Director.docx",
]
OUT = ROOT / "tmp" / "inventario_rev64"
OUT.mkdir(parents=True, exist_ok=True)

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def compact(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def paragraph_text(p) -> str:
    return compact("".join(t.text or "" for t in p.findall(".//w:t", NS)))


def xml_text(path: Path) -> tuple[list[str], dict]:
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    paragraphs = [paragraph_text(p) for p in root.findall(".//w:p", NS)]
    paragraphs = [p for p in paragraphs if p]
    stats = {
        "insertions": len(root.findall(".//w:ins", NS)),
        "deletions": len(root.findall(".//w:del", NS)),
        "comments_anchors": len(root.findall(".//w:commentRangeStart", NS)),
    }
    return paragraphs, stats


inventory = []
for path in FILES:
    doc = Document(path)
    xml_paragraphs, xml_stats = xml_text(path)
    headings = [
        {"style": p.style.name if p.style else "", "text": compact(p.text)}
        for p in doc.paragraphs
        if p.text.strip() and p.style and p.style.name.lower().startswith(("heading", "título", "titulo"))
    ]
    tables = []
    for idx, table in enumerate(doc.tables, 1):
        rows = [[compact(cell.text) for cell in row.cells] for row in table.rows]
        tables.append({"index": idx, "rows": rows})

    record = {
        "path": str(path),
        "size": path.stat().st_size,
        "paragraphs_python_docx": len(doc.paragraphs),
        "paragraphs_xml_nonempty": len(xml_paragraphs),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "sections": len(doc.sections),
        "headings": headings,
        **xml_stats,
    }
    inventory.append(record)
    (OUT / f"{path.stem}_text.txt").write_text("\n".join(xml_paragraphs), encoding="utf-8")
    (OUT / f"{path.stem}_tables.json").write_text(
        json.dumps(tables, ensure_ascii=False, indent=2), encoding="utf-8"
    )

(OUT / "inventory.json").write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(inventory, ensure_ascii=False, indent=2))
